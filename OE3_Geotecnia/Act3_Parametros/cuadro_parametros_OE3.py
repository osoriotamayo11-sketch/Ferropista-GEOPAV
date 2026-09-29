# -*- coding: utf-8 -*-
"""
OE 3 · Actividad 3 — Cuadro de parámetros geotécnicos adoptados y su rango de sensibilidad (sesión 17).

Salida: OE3_Geotecnia/Act3_Parametros/Cuadro_parametros_adoptados_OE3.xlsx
Seis juegos (decisión de Migue, 28 sep 2026): por portal, 1 de suelo superficial + 2 de roca.
  S1 suelo residual (portal oriental) · S2 ceniza volcánica (portal occidental)
  R1 PCAn neis · R2 Jcdi granodiorita · R3 Kqs · R4 Kqv
Todo valor lleva marca F / CP / H y su fuente. Las propiedades equivalentes del macizo (c', φ', σcm)
se calculan con FÓRMULAS de Excel (Hoek, Carranza-Torres y Corkum, 2002) a partir de las entradas.
Después de correr este script se recalcula con recalc.py (skill xlsx) o abriendo en Excel.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.comments import Comment

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.abspath(os.path.join(AQUI, "..", ".."))
SALIDA = os.path.join(AQUI, "Cuadro_parametros_adoptados_OE3.xlsx")
VERDE, AZUL, GRIS = "178E2C", "193F77", "475569"
F = "Arial"
fino = Side(style="thin", color="CBD5E1")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
ENTRADA = Font(name=F, size=10, color="0000FF")
NORMAL = Font(name=F, size=10)
NEGRA = Font(name=F, size=10, bold=True)
CAB = Font(name=F, size=10, bold=True, color="FFFFFF")
FILL_CAB = PatternFill("solid", fgColor=AZUL)
FILL_H = PatternFill("solid", fgColor="FFF7D6")

FUENTES = [
    ("CAR74", "Carrillo C., J. (1974). Propiedades físicas de los suelos derivados del Batolito antioqueño. Universidad Nacional de Colombia, p. 128.",
     "https://repositorio.unal.edu.co/items/a129a695-229f-4f86-b4f4-5cb217d198e1"),
    ("BBM13", "Betancur G., Y., Builes B., M. y Millán Á., Á. (2013). Variación de las propiedades mecánicas de arcillas alófanas en Colombia al variar el grado de saturación. Revista EIA 10(20), 173-181. Tabla 1 y figuras 2-3 (p. 179-180).",
     "http://www.scielo.org.co/pdf/eia/n20/n20a15.pdf"),
    ("HOEK", "Hoek, E. Practical Rock Engineering, cap. 11 «Rock mass properties». Tabla 2 (resistencia por campo) y tabla 3 (mi).",
     "https://www.rocscience.com/assets/resources/learning/hoek/Practical-Rock-Engineering-Chapter-11-Rock-Mass-Properties.pdf"),
    ("HCC02", "Hoek, E., Carranza-Torres, C. y Corkum, B. (2002). Hoek-Brown failure criterion – 2002 edition. Ecuaciones de mb, s, a, c′, φ′ y σcm (reproducidas en HOEK, cap. 11, ec. 11-17).",
     "https://www.rocscience.com/assets/resources/learning/hoek/Practical-Rock-Engineering-Chapter-11-Rock-Mass-Properties.pdf"),
    ("NSR10", "NSR-10, Título H, tabla H.4.7-1: factores de seguridad indirectos mínimos 3,0 / 2,5 / 1,5 (verificado en el texto de la norma).",
     "https://www.scg.org.co/Titulo-H-NSR-10-Decreto%20Final-2010-01-14.pdf"),
    ("SGC", "Planchas geológicas 244 Ibagué (1982) y 243 Armenia (1985), SGC/INGEOMINAS, leyendas impresas; memoria 244 y reseña 243. Ver producto de la Act 1.",
     "OE3_Geotecnia/Act1_Cartografia/Planos_geologicos_portales_OE3.pdf"),
]

# (clave, portal, unidad, material, [(parámetro, unidad, min, típ, máx, marca, fuente, nota)])
SUELOS = [
    ("S1", "Oriental (Ibagué)", "Suelo residual sobre PCAn / Jcdi", "Limo arenoso residual de rocas granitoides y neises (zona B del perfil de meteorización)", [
        ("c′ (cohesión efectiva)", "kPa", 0, 15, 29, "F", "CAR74",
         "Análogo: Batolito Antioqueño [H]. Zona B: c′ «alrededor de 0,3 kg/cm²» (≈29 kPa). Mínimo 0: el autor indica que la cohesión desaparece con saturación total."),
        ("φ′ (ángulo de fricción efectivo)", "°", 27, 30, 32, "F", "CAR74", "Zona B: φ′ entre 27° y 32°, esfuerzos efectivos."),
        ("γ (peso unitario total)", "kN/m³", 16, 17, 18, "H", "—",
         "Supuesto. CAR74 reporta «densidad natural» 1,28-1,32 g/cm³, que parece densidad seca; no se usa como γ total."),
    ]),
    ("S2", "Occidental (Calarcá)", "Cobertura de cenizas (Andisol) sobre Kqs / Kqv", "Ceniza volcánica alofánica (Eje Cafetero)", [
        ("c (cohesión, ensayo UU)", "kPa", 15, 25, 35, "F", "BBM13",
         "Análogo: Pereira [H]. Leído de la figura 2: 17-39 kPa para S = 71-97 %; decrece al saturarse. Ensayo no drenado, no efectivo."),
        ("φ (ángulo de fricción, UU)", "°", 23, 28, 33, "F", "BBM13", "Leído de la figura 3: 20-40° con dispersión; la mayoría entre 23° y 35°."),
        ("γ (peso unitario húmedo)", "kN/m³", 11.9, 13.5, 15.6, "F", "BBM13", "Tabla 1: 11,9-14,4 kN/m³ sin saturar; 12,6-15,6 saturadas. Típico: valor medio adoptado [H]."),
    ]),
]
# roca: (clave, portal, unidad, material, sci(min,tip,max), mi(min,tip,max), GSI(min,tip,max), nota_sci, nota_mi, nota_gsi)
ROCAS = [
    ("R1", "Oriental (Ibagué)", "PCAn — Neises y Anfibolitas de Tierradentro", "Neis cuarzo-feldespático-biotítico, meteorizado",
     (25, 60, 100), (23, 28, 33), (30, 42, 55),
     "Roca sana R5 (100-250 MPa, HOEK tabla 2). Se reduce por «intensa meteorización» (memoria 244) [H].",
     "Neis 28 ± 5 (HOEK tabla 3), medido normal a la foliación.",
     "Foliado y meteorizado: bloques/fracturado con superficies regulares a pobres [H]."),
    ("R2", "Oriental (Ibagué)", "Jcdi — Batolito de Ibagué", "Granodiorita biotítico-hornbléndica",
     (25, 80, 150), (26, 29, 32), (35, 48, 60),
     "Roca sana R5 (100-250 MPa, HOEK tabla 2). Reducida por meteorización [H].",
     "Granodiorita 29 ± 3 (HOEK tabla 3).",
     "Masiva a fracturada; superficies regulares [H]."),
    ("R3", "Occidental (Calarcá)", "Kqs — Fm. Quebradagrande, miembro sedimentario", "Grauvacas, lutitas, chert, calizas; localmente cataclasadas",
     (10, 30, 60), (6, 12, 18), (20, 30, 40),
     "Lutita R3-R4 (25-100 MPa); chert R6. Cataclasis (leyenda 243) y meteorización → rango bajo [H].",
     "Lutitas 6 ± 2 a grauvacas 18 ± 3 (HOEK tabla 3).",
     "Macizo heterogéneo tipo flysch, cataclasado (HOEK tabla 6) [H]."),
    ("R4", "Occidental (Calarcá)", "Kqv — Fm. Quebradagrande, miembro volcánico", "Diabasas y andesitas; sectores de intensa cataclasis",
     (25, 70, 150), (15, 20, 25), (25, 38, 50),
     "Diabasa R6 (> 250 MPa) en roca sana; intensa cataclasis (leyenda 243) → rango reducido [H].",
     "Diabasa 15 ± 5 a andesita 25 ± 5 (HOEK tabla 3).",
     "Fracturado por cataclasis [H]."),
]

wb = Workbook()
# ------------------------------------------------------------------ Portada
p = wb.active; p.title = "Portada"
p.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", (3, 22, 22, 22, 22, 22, 22, 3)):
    p.column_dimensions[col].width = w
for nombre, celda, alto in (("logo-unibague.png", "B2", 70), ("logo-geopav.png", "F2", 70)):
    ruta = os.path.join(RAIZ, "public", nombre)
    if os.path.exists(ruta):
        im = XLImage(ruta); r = alto / im.height; im.height = alto; im.width = int(im.width * r); p.add_image(im, celda)
p["B8"] = "Cuadro de parámetros geotécnicos adoptados y su rango de sensibilidad"
p["B8"].font = Font(name=F, size=18, bold=True, color=VERDE)
p["B9"] = "Objetivo específico 3 · Actividad 3 — Geotecnia de los portales del túnel Ferropista"
p["B9"].font = Font(name=F, size=12, color=AZUL)
ficha = [
    ("Objetivo", "OE 3: Estimar la capacidad portante preliminar del terreno de fundación de las infraestructuras de superficie en los portales de Ibagué y Armenia"),
    ("Actividad", "3. Definición de los parámetros geotécnicos de diseño adoptados y de su rango de sensibilidad"),
    ("Entregable / formato", "Cuadro de parámetros adoptados · Documento Excel"),
    ("Periodo", "Bloque 2 del Plan de Acción (5 oct - 7 nov 2026)"),
    ("Responsable asignado", "Tamayo Osorio Miguel Ángel"),
    ("Semillero", "Semillero de Investigación GEOPAV · Ingeniería Civil · Universidad de Ibagué · Paz y Región 2026B"),
    ("Fecha de elaboración", "28 sep 2026"),
    ("Advertencia", "NINGÚN valor proviene de exploración en los portales. Son rangos de fuentes publicadas de materiales análogos o supuestos declarados. No son parámetros de diseño: sirven para el análisis de sensibilidad de la actividad 4 y para dimensionar la exploración que haría falta."),
]
for i, (k, v) in enumerate(ficha):
    r = 11 + i
    p.cell(r, 2, k).font = NEGRA; p.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)
    p.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    c = p.cell(r, 3, v); c.font = NORMAL if k != "Advertencia" else Font(name=F, size=10, bold=True, color="9A3412")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    p.row_dimensions[r].height = 30 if len(v) < 110 else 58
r = 11 + len(ficha) + 1
p.cell(r, 2, "Índice de hojas").font = Font(name=F, size=12, bold=True, color=AZUL); r += 1
indice = [("Resumen", "Tabla 1. Resumen de los seis juegos de parámetros (valor típico y rango)"),
          ("Suelos", "Tabla 2. Suelos superficiales S1 y S2 (entradas de la actividad 4, Meyerhof)"),
          ("Roca", "Tabla 3. Roca R1-R4: entradas Hoek-Brown y propiedades equivalentes del macizo [CP]"),
          ("Marcas y fuentes", "Tabla 4. Marcas F / CP / H y Tabla 5. Fuentes citadas")]
for h, d in indice:
    p.cell(r, 2, h).font = NEGRA; p.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7); p.cell(r, 3, d).font = NORMAL; r += 1
r += 1
p.merge_cells(start_row=r, start_column=2, end_row=r + 2, end_column=7)
c = p.cell(r, 2, "Índice de anexos: este cuadro no tiene anexos. Índice de figuras: no tiene figuras. "
            "Elaborado con apoyo de IA (Claude) y verificado contra las fuentes citadas: cada valor se leyó en el documento original descargado. "
            "Script: OE3_Geotecnia/Act3_Parametros/cuadro_parametros_OE3.py")
c.font = Font(name=F, size=9, italic=True, color=GRIS); c.alignment = Alignment(wrap_text=True, vertical="top")

def cabecera(ws, fila, textos, anchos=None):
    for j, t in enumerate(textos, start=1):
        c = ws.cell(fila, j, t); c.font = CAB; c.fill = FILL_CAB; c.border = BORDE
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    if anchos:
        for j, w in enumerate(anchos, start=1):
            ws.column_dimensions[chr(64 + j)].width = w

def titulo(ws, texto, fila=1):
    ws[f"A{fila}"] = texto; ws[f"A{fila}"].font = Font(name=F, size=13, bold=True, color=AZUL)

def celda(ws, r, c, v, fuente=NORMAL, fmt=None, marca=None, wrap=False):
    x = ws.cell(r, c, v); x.font = fuente; x.border = BORDE
    x.alignment = Alignment(wrap_text=wrap, vertical="top")
    if fmt: x.number_format = fmt
    if marca == "H": x.fill = FILL_H
    return x

# ------------------------------------------------------------------ Suelos
s = wb.create_sheet("Suelos"); s.sheet_view.showGridLines = False
titulo(s, "Tabla 2. Suelos superficiales: parámetros adoptados y rango de sensibilidad")
s["A2"] = "Texto azul = entrada editable. Fondo amarillo = hipótesis [H]. Estos juegos alimentan la capacidad portante admisible (actividad 4, Meyerhof)."
s["A2"].font = Font(name=F, size=9, italic=True, color=GRIS)
cabecera(s, 4, ["Juego", "Portal", "Material", "Parámetro", "Unidad", "Mínimo", "Típico", "Máximo", "Marca", "Fuente", "Nota"],
         [7, 17, 30, 28, 8, 9, 9, 9, 7, 9, 70])
r = 5; celdas_suelo = {}
for clave, portal, unidad, material, pars in SUELOS:
    for (par, u, mn, tp, mx, marca, fte, nota) in pars:
        celda(s, r, 1, clave, NEGRA); celda(s, r, 2, portal); celda(s, r, 3, f"{unidad}: {material}", wrap=True)
        celda(s, r, 4, par); celda(s, r, 5, u)
        for j, v in zip((6, 7, 8), (mn, tp, mx)):
            celda(s, r, j, v, ENTRADA, "0.0" if isinstance(v, float) else "0", marca)
        celda(s, r, 9, marca, NEGRA); celda(s, r, 10, fte); celda(s, r, 11, nota, wrap=True)
        s.row_dimensions[r].height = 42
        celdas_suelo[(clave, par.split(" ")[0])] = r
        r += 1
s.cell(r + 1, 1, "Fuente: CAR74 y BBM13 (ver hoja «Marcas y fuentes»). Los análogos no son del sitio del portal [H].").font = Font(name=F, size=9, italic=True, color=GRIS)
s.freeze_panes = "A5"

# ------------------------------------------------------------------ Roca
k = wb.create_sheet("Roca"); k.sheet_view.showGridLines = False
titulo(k, "Tabla 3. Roca: entradas Hoek-Brown y propiedades equivalentes del macizo rocoso")
k["A2"] = ("Criterio generalizado de Hoek-Brown (HCC02). Entradas en azul; propiedades equivalentes calculadas con fórmulas [CP] "
           "para el esfuerzo de confinamiento máximo σ3max de la celda C3 (zona de influencia de una cimentación superficial, supuesto [H]).")
k["A2"].font = Font(name=F, size=9, italic=True, color=GRIS)
k["A3"] = "σ3max (MPa) [H]"; k["A3"].font = NEGRA
k["C3"] = 0.5; k["C3"].font = ENTRADA; k["C3"].fill = FILL_H; k["C3"].border = BORDE
k["C3"].comment = Comment("Supuesto [H]: confinamiento máximo en la zona de influencia de una zapata de pocos metros. Cambiarlo recalcula c′ y φ′.", "GEOPAV")
k["E3"] = "D (factor de alteración) [H]"; k["E3"].font = NEGRA
k["G3"] = 0; k["G3"].font = ENTRADA; k["G3"].fill = FILL_H; k["G3"].border = BORDE
k["G3"].comment = Comment("D = 0: excavación sin voladura de daño (HOEK tabla 7). Supuesto [H].", "GEOPAV")
hdr = ["Juego", "Portal", "Unidad", "Caso", "σci (MPa)", "mi", "GSI", "mb", "s", "a", "σ3n", "φ′ (°)", "c′ (kPa)", "σcm (MPa)", "γ (kN/m³) [H]"]
cabecera(k, 5, hdr, [7, 17, 36, 9, 10, 8, 8, 9, 11, 9, 9, 9, 10, 10, 12])
r = 6
for clave, portal, unidad, material, sci, mi, gsi, nsci, nmi, ngsi in ROCAS:
    for idx, caso in enumerate(("Mínimo", "Típico", "Máximo")):
        celda(k, r, 1, clave, NEGRA); celda(k, r, 2, portal); celda(k, r, 3, f"{unidad} — {material}", wrap=True); celda(k, r, 4, caso)
        celda(k, r, 5, sci[idx], ENTRADA, "0", "H"); celda(k, r, 6, mi[idx], ENTRADA, "0"); celda(k, r, 7, gsi[idx], ENTRADA, "0", "H")
        E, Fm, G = f"E{r}", f"F{r}", f"G{r}"
        celda(k, r, 8, f"={Fm}*EXP(({G}-100)/(28-14*$G$3))", NORMAL, "0.000")
        celda(k, r, 9, f"=EXP(({G}-100)/(9-3*$G$3))", NORMAL, "0.00000")
        celda(k, r, 10, f"=0.5+(EXP(-{G}/15)-EXP(-20/3))/6", NORMAL, "0.000")
        celda(k, r, 11, f"=$C$3/{E}", NORMAL, "0.0000")
        H_, I_, J_, K_ = f"H{r}", f"I{r}", f"J{r}", f"K{r}"
        t = f"(6*{J_}*{H_}*({I_}+{H_}*{K_})^({J_}-1))"
        celda(k, r, 12, f"=DEGREES(ASIN({t}/(2*(1+{J_})*(2+{J_})+{t})))", NORMAL, "0.0")
        celda(k, r, 13, f"=1000*{E}*((1+2*{J_})*{I_}+(1-{J_})*{H_}*{K_})*({I_}+{H_}*{K_})^({J_}-1)"
                        f"/((1+{J_})*(2+{J_})*SQRT(1+{t}/((1+{J_})*(2+{J_}))))", NORMAL, "0")
        celda(k, r, 14, f"={E}*({H_}+4*{I_}-{J_}*({H_}-8*{I_}))*({H_}/4+{I_})^({J_}-1)/(2*(1+{J_})*(2+{J_}))", NORMAL, "0.00")
        celda(k, r, 15, (25, 26, 27)[idx], ENTRADA, "0", "H")
        k.row_dimensions[r].height = 30
        r += 1
r += 1
k.cell(r, 1, "Notas por juego").font = Font(name=F, size=11, bold=True, color=AZUL); r += 1
for clave, portal, unidad, material, sci, mi, gsi, nsci, nmi, ngsi in ROCAS:
    for etiqueta, texto, marca in (("σci", nsci, "H"), ("mi", nmi, "F"), ("GSI", ngsi, "H")):
        celda(k, r, 1, clave, NEGRA); celda(k, r, 2, etiqueta); celda(k, r, 3, f"[{marca}] {texto}", wrap=True)
        k.merge_cells(start_row=r, start_column=3, end_row=r, end_column=15); k.row_dimensions[r].height = 18; r += 1
k.cell(r + 1, 1, "Fuente: HOEK (tablas 2 y 3), HCC02 (ecuaciones). γ de roca: rango típico asumido [H], sin tabla de fuente verificada.").font = Font(name=F, size=9, italic=True, color=GRIS)
k.cell(r + 2, 1, "Lectura: con σ3max bajo (cimentación superficial) el ajuste Mohr-Coulomb de Hoek-Brown da φ′ altos y c′ moderados; describen el macizo, NO son parámetros para Meyerhof, que se aplica a los suelos S1 y S2.").font = Font(name=F, size=9, italic=True, color="9A3412")
k.freeze_panes = "A6"

# ------------------------------------------------------------------ Resumen
z = wb.create_sheet("Resumen", 1); z.sheet_view.showGridLines = False
titulo(z, "Tabla 1. Resumen de los seis juegos de parámetros (valor típico y rango mínimo-máximo)")
cabecera(z, 3, ["Juego", "Portal", "Tipo", "Unidad / material", "c′ típ. (kPa)", "c′ rango", "φ′ típ. (°)", "φ′ rango", "γ típ. (kN/m³)", "Uso"],
         [7, 17, 8, 44, 11, 13, 10, 12, 12, 34])
r = 4
for clave, portal, unidad, material, pars in SUELOS:
    rc = celdas_suelo[(clave, "c′")] if (clave, "c′") in celdas_suelo else celdas_suelo[(clave, "c")]
    rf = celdas_suelo[(clave, "φ′")] if (clave, "φ′") in celdas_suelo else celdas_suelo[(clave, "φ")]
    rg = celdas_suelo[(clave, "γ")]
    vals = [clave, portal, "Suelo", f"{unidad}: {material}",
            f"=Suelos!G{rc}", f'=TEXT(Suelos!F{rc},"0")&" – "&TEXT(Suelos!H{rc},"0")',
            f"=Suelos!G{rf}", f'=TEXT(Suelos!F{rf},"0")&" – "&TEXT(Suelos!H{rf},"0")',
            f"=Suelos!G{rg}", "Capacidad portante admisible (Act 4, Meyerhof)"]
    for j, v in enumerate(vals, start=1):
        celda(z, r, j, v, NORMAL, "0.0" if j == 9 else ("0" if j in (5, 7) else None), wrap=(j in (4, 10)))
    z.row_dimensions[r].height = 30; r += 1
fila_roca = 6
for clave, portal, unidad, material, *_ in ROCAS:
    a, b, c_ = fila_roca, fila_roca + 1, fila_roca + 2
    vals = [clave, portal, "Roca", f"{unidad} — {material}",
            f"=Roca!M{b}", f'=TEXT(Roca!M{a},"0")&" – "&TEXT(Roca!M{c_},"0")',
            f"=Roca!L{b}", f'=TEXT(Roca!L{a},"0")&" – "&TEXT(Roca!L{c_},"0")',
            f"=Roca!O{b}", "Referencia (roca bajo la cobertura); no entra en Meyerhof"]
    for j, v in enumerate(vals, start=1):
        celda(z, r, j, v, NORMAL, "0.0" if j == 9 else ("0" if j in (5, 7) else None), wrap=(j in (4, 10)))
    z.row_dimensions[r].height = 30; r += 1; fila_roca += 3
z.cell(r + 1, 1, "Fuente: hojas «Suelos» y «Roca». En roca, c′ y φ′ son equivalentes Mohr-Coulomb del macizo [CP] para σ3max de Roca!C3.").font = Font(name=F, size=9, italic=True, color=GRIS)
z.cell(r + 2, 1, "Factores de seguridad para la actividad 4 (NSR-10, tabla H.4.7-1, verificada): CM+CV normal 3,0 · CM+CV máxima 2,5 · CM+CV normal+sismo seudoestático 1,5.").font = Font(name=F, size=9, italic=True, color=GRIS)

# ------------------------------------------------------------------ Marcas y fuentes
m = wb.create_sheet("Marcas y fuentes"); m.sheet_view.showGridLines = False
titulo(m, "Tabla 4. Marcas de origen")
cabecera(m, 3, ["Marca", "Significado"], [12, 110])
for i, (a, b) in enumerate((("F", "Dato de fuente publicada, leído en el documento original descargado"),
                            ("CP", "Cálculo propio, reproducible con las fórmulas de este libro y el script citado"),
                            ("H", "Hipótesis o supuesto del semillero, declarado; fondo amarillo en las celdas"))):
    celda(m, 4 + i, 1, a, NEGRA); celda(m, 4 + i, 2, b)
m["A9"] = "Tabla 5. Fuentes citadas"; m["A9"].font = Font(name=F, size=13, bold=True, color=AZUL)
cabecera(m, 10, ["Clave", "Referencia", "Enlace / ruta"], [12, 110, 70])
m.column_dimensions["C"].width = 70
for i, (a, b, c_) in enumerate(FUENTES):
    celda(m, 11 + i, 1, a, NEGRA); celda(m, 11 + i, 2, b, wrap=True); celda(m, 11 + i, 3, c_, wrap=True)
    m.row_dimensions[11 + i].height = 32

wb.save(SALIDA)
print("escrito", SALIDA)
