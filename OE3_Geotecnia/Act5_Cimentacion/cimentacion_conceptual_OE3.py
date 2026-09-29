# -*- coding: utf-8 -*-
"""
OE 3 · Actividad 5 — Diseño conceptual de la cimentación superficial de las terminales Ro-Ro (sesión 18).

Paso 1 de 3: CÁLCULO. Produce cimentacion_conceptual_OE3.json, que consumen
  plano_cimentacion_OE3.py      → plano .dxf (AutoCAD) + lámina PDF/PNG
  informe_cimentacion_OE3.py    → informe PDF con el estándar del semillero

Encadenamiento (ninguna cifra geotécnica se escribe a mano):
  - Parámetros del suelo, hipótesis y FS: leídos de la hoja «Entradas» de la memoria de la Act 4
    (Memoria_capacidad_portante_OE3.xlsx, valores ya recalculados).
  - Capacidad admisible qadm(B): mismas ecuaciones que la Act 4 (Meyerhof + FSB directo + asentamiento),
    importadas de memoria_capacidad_portante_OE3.py. CONTROL: antes de dimensionar, se reproduce la columna
    «qadm NSR-10 H.4.2.3» de la Act 4 para B = 1,5 / 2,0 / 3,0 m; si difiere más de 0,1 %, el script se detiene.
Cargas: NSR-10 Título B (tablas B.3.2-1, B.3.4.1-4, B.3.4.3-1, B.4.2.1-1, B.4.2.1-2; combinación B.2.3-2/-3,
verificadas en el texto de la norma) y Resolución 4100 de 2004 del Ministerio de Transporte (arts. 8 y 9: 3S3,
52.000 kg; trídem de 12 llantas, 24.000 kg). Geometría de las estructuras: hipótesis [H] declaradas.
"""
import json
import math
import os
import sys

from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
OE3 = os.path.dirname(AQUI)
RAIZ = os.path.dirname(OE3)
ACT4 = os.path.join(OE3, "Act4_Capacidad")
sys.path.insert(0, ACT4)
from memoria_capacidad_portante_OE3 import meyerhof_qult  # noqa: E402  (misma ecuación que la Act 4)

MEMORIA4 = os.path.join(ACT4, "Memoria_capacidad_portante_OE3.xlsx")
SGC = os.path.join(OE3, "Act1_Cartografia", "consulta_sgc_portales_OE3.json")
PEND = os.path.join(OE3, "Act2_Amenaza", "pendiente_portales_OE3.json")
SALIDA = os.path.join(AQUI, "cimentacion_conceptual_OE3.json")

G = 9.81
GAMMA_C = 2400 * G / 1000          # kN/m³ — concreto reforzado 2.400 kg/m³ (NSR-10 tabla B.3.2-1) [F→CP]
GAMMA_RELLENO = 20.0               # kN/m³ — peso medio zapata + relleno sobre ella [H]
PASO_B, B_MIN, B_MAX = 0.1, 1.0, 6.0

# ------------------------------------------------------------------ hipótesis de geometría y cargas
H = {
    "anden_longitud_m": (750, "F", "Trenes de hasta 750 m (ponencia, dia. 23)"),
    "marquesina_luz_m": (12.0, "H", "Cubre la vía y un andén peatonal de servicio; pórtico metálico de una luz"),
    "marquesina_separacion_m": (10.0, "H", "Separación entre pórticos a lo largo del andén"),
    "marquesina_altura_libre_m": (7.0, "H", "Mayor que la altura máxima de un vehículo de carga, 4,40 m (Res. 4100/2004, art. 7), más la plataforma del vagón"),
    "cubierta_tablero_kPa": (0.08, "F", "Tablero metálico calibre 20 (NSR-10 tabla B.3.4.1-4)"),
    "cubierta_estructura_kPa": (0.25, "H", "Peso propio de correas y cerchas metálicas por m² en planta"),
    "columna_metalica_kN_m": (0.5, "H", "Peso lineal de la columna metálica"),
    "cubierta_Lr_kPa": (0.50, "F", "Cubierta metálica con pendiente ≤ 15° (NSR-10 tabla B.4.2.1-2)"),
    "edificio_luz_m": (6.0, "H", "Pórticos de concreto de 6 × 6 m; planta de 18 × 12 m (3 × 2 luces), 2 pisos"),
    "edificio_planta_m": ((18.0, 12.0), "H", "Edificio de control y operación de la terminal"),
    "edificio_altura_piso_m": (2.9, "H", "Menor de 3 m: permite usar la tabla B.3.4.3-1"),
    "losa_espesor_m": (0.15, "H", "Losa maciza equivalente (peso de losa aligerada)"),
    "oficinas_particiones_kPa": (2.0, "F", "Oficinas, particiones fijas de mampostería (NSR-10 tabla B.3.4.3-1)"),
    "oficinas_afinado_kPa": (1.8, "F", "Afinado de piso y cubierta, oficinas (NSR-10 tabla B.3.4.3-1)"),
    "oficinas_L_kPa": (2.0, "F", "Oficinas (NSR-10 tabla B.4.2.1-1); cubierta con la misma carga (nota 1 de la tabla B.4.2.1-2)"),
    "viga_seccion_m": ((0.40, 0.35), "H", "Viga de 0,40 × 0,50 m; 0,35 m bajo la losa"),
    "columna_seccion_m": ((0.40, 0.40), "H", "Columna de concreto"),
    "tridem_kN": (24000 * G / 1000, "F", "Eje trídem de 12 llantas, 24.000 kg (Res. 4100/2004, art. 9)"),
    "tractomula_3S3_kN": (52000 * G / 1000, "F", "Tracto-camión 3S3, 52.000 kg (Res. 4100/2004, art. 8)"),
    "tridem_huella_m": ((2.5, 2.8), "H", "Huella del trídem: ancho de ejes × longitud de los tres ejes"),
    "factor_dinamico": (1.30, "H", "Incremento por impacto del vehículo en rampa (sin tabla verificada en la NSR-10 para camiones)"),
    "rampa_losa_m": (0.25, "H", "Losa de concreto de la rampa"),
    "rampa_base_m": (0.30, "H", "Base granular compactada bajo la losa"),
    "rampa_altura_m": (1.0, "H", "Desnivel entre la vía de acceso y la plataforma del vagón"),
    "rampa_pendiente": (0.08, "H", "Pendiente longitudinal de la rampa"),
    "rampa_ancho_m": (4.0, "H", "Un carril de 3,65 m más bordillos"),
    "base_granular_kN_m3": (20.0, "H", "Peso unitario de la base granular"),
    "Df_m": (None, "H", "Profundidad de desplante: la misma de la Act 4"),
}
V = {k: v[0] for k, v in H.items()}


def leer_act4():
    wb = load_workbook(MEMORIA4, data_only=True)
    e = wb["Entradas"]
    par = {}
    for r in range(5, 40):
        j, p = e.cell(r, 1).value, e.cell(r, 3).value
        if j in ("S1", "S2") and p:
            clave = ("c" if p.startswith("c") else "phi" if p.startswith("φ") else "gamma" if p.startswith("γ")
                     else "E" if p.startswith("E") else "nu" if p.startswith("ν") else None)
            if clave:
                par.setdefault(j, {})[clave] = tuple(float(e.cell(r, c).value) for c in (5, 6, 7))
    hip = {}
    for r in range(15, 40):
        n = e.cell(r, 2).value
        if isinstance(n, str) and e.cell(r, 3).value is not None and isinstance(e.cell(r, 3).value, (int, float)):
            hip[n.split(",")[0]] = float(e.cell(r, 3).value)
            if n.startswith("CM + CV normal") or n.startswith("CM + CV"):
                pass
    fs = {}
    for r in range(15, 45):
        n = e.cell(r, 2).value
        if isinstance(n, str) and n.startswith("CM + CV"):
            fs[n] = (float(e.cell(r, 3).value), float(e.cell(r, 4).value))
    calc = wb["Cálculo"]
    ref = {}
    for r in range(5, 41):
        ref[(calc.cell(r, 1).value, calc.cell(r, 3).value, float(calc.cell(r, 4).value), calc.cell(r, 5).value)] = float(calc.cell(r, 25).value)
    return par, hip, fs, ref


PAR, HIP, FS, REF = leer_act4()
DF, GW, RHO, CD = HIP["Df"], HIP["γw"], HIP["ρadm"], HIP["Cd"]
V["Df_m"] = DF
FSI, FSB = FS["CM + CV normal"]
CASOS = ("Mínimo", "Típico", "Máximo")
NF = ("NF en la base", "NF profundo")


def reducir(c, phi, F):
    return c / F, math.degrees(math.atan(math.tan(math.radians(phi)) / F))


def qadm(j, caso, B, nf, cd=None):
    i = CASOS.index(caso)
    c, phi, g = (PAR[j][k][i] for k in ("c", "phi", "gamma"))
    E, nu = PAR[j]["E"][i], PAR[j]["nu"][i]
    ge = g - GW if nf == "NF en la base" else g
    qu = meyerhof_qult(c, phi, g, ge, B, DF)
    qf = min(qu / FSI, meyerhof_qult(*reducir(c, phi, FSB), g, ge, B, DF))
    qs = RHO * E / ((cd or CD) * B * (1 - nu ** 2)) + g * DF
    return min(qf, qs), ("Asentamiento" if qs < qf else "Falla"), qu


# ------------------------------------------------------------------ control de encadenamiento con la Act 4
dif = max(abs(qadm(j, c, B, nf)[0] - v) / v for (j, c, B, nf), v in REF.items())
if dif > 1e-3:
    raise SystemExit(f"DETENIDO: qadm no reproduce la Act 4 (diferencia relativa {dif:.2e}).")
print(f"control Act 4 OK: 36 casos reproducidos (dif. máx. {dif:.1e})")

# ------------------------------------------------------------------ cargas por apoyo [CP]
luz, sep = V["marquesina_luz_m"], V["marquesina_separacion_m"]
a_m = sep * luz / 2
PD_m = a_m * (V["cubierta_tablero_kPa"] + V["cubierta_estructura_kPa"]) + V["columna_metalica_kN_m"] * V["marquesina_altura_libre_m"]
PL_m = a_m * V["cubierta_Lr_kPa"]

l_e = V["edificio_luz_m"]
a_e = l_e * l_e
losa = V["losa_espesor_m"] * GAMMA_C
bv, hv = V["viga_seccion_m"]
bc, hc = V["columna_seccion_m"]
vigas = 2 * l_e * bv * hv * GAMMA_C
col = bc * hc * V["edificio_altura_piso_m"] * GAMMA_C
D_piso = a_e * (losa + V["oficinas_particiones_kPa"] + V["oficinas_afinado_kPa"]) + vigas + col
D_cub = a_e * (losa + V["oficinas_afinado_kPa"]) + vigas + col
PD_e, PL_e = D_piso + D_cub, 2 * a_e * V["oficinas_L_kPa"]

CARGAS = {
    "Z-1": dict(estructura="Marquesina metálica del andén de carga (columna típica)", area_aferente_m2=a_m,
                PD_kN=PD_m, PL_kN=PL_m, P_kN=PD_m + PL_m, combinacion="D + Lr (NSR-10 ec. B.2.3-3)", h_m=0.40),
    "Z-2": dict(estructura="Edificio de control, 2 pisos, pórticos de concreto (columna interior)", area_aferente_m2=a_e,
                PD_kN=PD_e, PL_kN=PL_e, P_kN=PD_e + PL_e, combinacion="D + L (NSR-10 ec. B.2.3-2)", h_m=0.60),
}


def dimensionar(P, j, caso, nf):
    B = B_MIN
    while B <= B_MAX + 1e-9:
        q_act = P / B ** 2 + GAMMA_RELLENO * DF
        qa, gob, _ = qadm(j, caso, B, nf)
        if q_act <= qa:
            return dict(B_m=round(B, 2), q_act_kPa=q_act, qadm_kPa=qa, gobierna=gob, uso=q_act / qa)
        B = round(B + PASO_B, 2)
    qa, gob, _ = qadm(j, caso, B_MAX, nf)
    return dict(B_m=None, q_act_kPa=P / B_MAX ** 2 + GAMMA_RELLENO * DF, qadm_kPa=qa, gobierna=gob, uso=None)


PORTAL = {"S1": "Oriental (Ibagué)", "S2": "Occidental (Calarcá)"}
zapatas = []
for z, cg in CARGAS.items():
    for j in ("S1", "S2"):
        for caso in CASOS:
            for nf in NF:
                d = dimensionar(cg["P_kN"], j, caso, nf)
                zapatas.append(dict(zapata=z, suelo=j, portal=PORTAL[j], caso=caso, nf=nf, **d))

# ------------------------------------------------------------------ losa de cimentación del edificio (alternativa)
Lx, Ly = V["edificio_planta_m"]
n_col = (Lx / l_e + 1) * (Ly / l_e + 1)
# carga total ≈ columnas interiores equivalentes: área total × carga por m² de la columna interior
w_D = PD_e / a_e
w_L = PL_e / a_e
losa_fund_m = 0.40  # [H]
q_losa = w_D + w_L + losa_fund_m * GAMMA_C
CD_RECT = 1.15   # FHWA 089 tabla 8-13, rectángulo L/B = 1,5, promedio (flexible)
losas = []
for j in ("S1", "S2"):
    for caso in CASOS:
        qa, gob, _ = qadm(j, caso, Ly, "NF en la base", cd=CD_RECT)
        losas.append(dict(suelo=j, portal=PORTAL[j], caso=caso, q_act_kPa=q_losa, qadm_kPa=qa, gobierna=gob, cumple=q_losa <= qa))

# ------------------------------------------------------------------ rampa de acceso de tractomulas
bx, by = V["tridem_huella_m"]
z_ = V["rampa_losa_m"] + V["rampa_base_m"]
area_z = (bx + z_) * (by + z_)                   # repartición 2:1 hasta la subrasante [H]
q_trid = V["factor_dinamico"] * V["tridem_kN"] / area_z
q_pp = V["rampa_losa_m"] * GAMMA_C + V["rampa_base_m"] * V["base_granular_kN_m3"]
q_rampa = q_trid + q_pp
L_rampa = V["rampa_altura_m"] / V["rampa_pendiente"]
rampas = []
for j in ("S1", "S2"):
    for caso in CASOS:
        B_eq = min(bx + z_, by + z_)
        # la huella cargada se trata como zapata superficial (Df ≈ espesor del paquete) — [H]
        i = CASOS.index(caso)
        c, phi, g = (PAR[j][k][i] for k in ("c", "phi", "gamma"))
        E, nu = PAR[j]["E"][i], PAR[j]["nu"][i]
        qu = meyerhof_qult(c, phi, g, g - GW, B_eq, z_)
        qf = min(qu / FSI, meyerhof_qult(*reducir(c, phi, FSB), g, g - GW, B_eq, z_))
        qs = RHO * E / (CD * B_eq * (1 - nu ** 2)) + g * z_
        qa = min(qf, qs)
        rampas.append(dict(suelo=j, portal=PORTAL[j], caso=caso, B_eq_m=B_eq, q_act_kPa=q_rampa, qadm_kPa=qa,
                           gobierna="Asentamiento" if qs < qf else "Falla", cumple=q_rampa <= qa))

sgc = json.load(open(SGC, encoding="utf-8"))["portales"]
pend = json.load(open(PEND, encoding="utf-8"))["portales"]
sismo = {k: dict(Aa=v["nsr10"][0]["AA"], Av=v["nsr10"][0]["AV"], zona=v["nsr10"][0]["ZONA_AMENAZA_SÍSMICA"]) for k, v in sgc.items()}
pendiente = {k: dict(media_250m=v["circulos"]["250"]["media"], p90_250m=v["circulos"]["250"]["p90"]) for k, v in pend.items()}

salida = dict(
    generado_por="OE3_Geotecnia/Act5_Cimentacion/cimentacion_conceptual_OE3.py",
    control_act4_dif_relativa=dif,
    hipotesis={k: dict(valor=v[0] if k != "Df_m" else DF, marca=v[1], justificacion=v[2]) for k, v in H.items()},
    constantes=dict(gamma_concreto_kN_m3=GAMMA_C, gamma_relleno_kN_m3=GAMMA_RELLENO, Df_m=DF, rho_adm_m=RHO, Cd=CD,
                    FSICP=FSI, FSBM=FSB, paso_B_m=PASO_B, B_min_m=B_MIN, B_max_m=B_MAX),
    cargas=CARGAS, zapatas=zapatas,
    losa_edificio=dict(espesor_m=losa_fund_m, planta_m=[Lx, Ly], n_columnas=n_col, q_uniforme_kPa=q_losa, Cd=CD_RECT, casos=losas),
    rampa=dict(q_tridem_kPa=q_trid, q_peso_propio_kPa=q_pp, q_total_kPa=q_rampa, area_reparticion_m2=area_z,
               longitud_m=L_rampa, casos=rampas),
    sismo_SGC=sismo, pendiente_Act2=pendiente,
)
json.dump(salida, open(SALIDA, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("escrito", SALIDA)
for z in zapatas:
    print(z["zapata"], z["suelo"], z["caso"], z["nf"], z["B_m"], round(z["q_act_kPa"]), round(z["qadm_kPa"]), z["gobierna"])
print("cargas", {k: round(v["P_kN"], 1) for k, v in CARGAS.items()})
print("losa", round(q_losa, 1), [(l["suelo"], l["caso"], round(l["qadm_kPa"]), l["cumple"]) for l in losas])
print("rampa", round(q_rampa, 1), round(L_rampa, 1), [(r["suelo"], r["caso"], round(r["qadm_kPa"]), r["cumple"]) for r in rampas])
