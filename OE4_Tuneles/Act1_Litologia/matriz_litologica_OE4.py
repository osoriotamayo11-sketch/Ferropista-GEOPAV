# -*- coding: utf-8 -*-
"""
OE 4 · Actividad 1 — Caracterización litológica y de zonas de falla del macizo a partir de la cartografía del SGC.
Producto: «Matriz litológica» · Documento Excel. Sesión 18 (28-29 sep 2026).

Qué hace
  1. Consulta los servicios ArcGIS REST del Servicio Geológico Colombiano a lo largo del eje del túnel (OE 1):
       - Mapa Geológico de Colombia 2023 (1:1.000.000): unidades (capa 733) y fallas (capa 704)      ← FUENTE PRINCIPAL
       - Atlas Geológico de Colombia 2007 (1:500.000): unidades (13) y fallas (11)                      ← contraste
       - Geología del departamento del Tolima: unidades (188) y fallas (187)                             ← contraste
       - Geología del departamento del Quindío (servicio de geomecánica): unidades (19)                  ← contraste
     Guarda la respuesta cruda en Datos/ (GeoJSON) con fecha; con --cache reutiliza esa copia sin red.
  2. Proyecta cada intersección sobre el eje (EPSG 3116) y sectoriza por abscisa (PK, m desde el portal oriental).
  3. Cruza cada tramo con la cobertura del perfil del OE 1 (perfil_continuo.csv) [CP].
  4. Asigna a cada tramo un juego de parámetros de macizo (σci, mi, GSI mín/típ/máx). Donde el tramo coincide con
     una unidad ya parametrizada en la Act 3 del OE 3, LEE esos valores del Excel de la Act 3 (sin copiarlos a mano).
  5. Escribe matriz_litologica_OE4.json (lo consumen las Act 2-5) y Matriz_litologica_OE4.xlsx.
Análogo medido: Túnel de La Línea (piloto y II Centenario), cuyo portal oriental está a 18 m en planta del eje
(PK ≈ 37,2 km) y que se abre a 30° hasta 4,4 km en el PK 44,6. Fuentes: Castro y Pérez (2013) tabla 2 y figura 4;
Dávila (2015) tabla de fallas. Distancias calculadas con OE2_Plataforma/Act2_Visor/via_actual_ruta40.geojson [CP].
Requiere shapely, pyproj, openpyxl.
"""
import csv
import datetime
import hashlib
import json
import math
import os
import sys
import urllib.parse
import urllib.request

import pyproj
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from shapely.geometry import LineString, Point, shape
from shapely.ops import transform

AQUI = os.path.dirname(os.path.abspath(__file__))
OE4 = os.path.dirname(AQUI)
RAIZ = os.path.dirname(OE4)
DATOS = os.path.join(AQUI, "Datos")
os.makedirs(DATOS, exist_ok=True)
TRAZADO = os.path.join(RAIZ, "OE1_Topografia", "Act4_Trazado", "trazado_tunel.geojson")
PERFIL = os.path.join(RAIZ, "OE1_Topografia", "Act5_Perfil", "perfil_continuo.csv")
VIA = os.path.join(RAIZ, "OE2_Plataforma", "Act2_Visor", "via_actual_ruta40.geojson")
ACT3 = os.path.join(RAIZ, "OE3_Geotecnia", "Act3_Parametros", "Cuadro_parametros_adoptados_OE3.xlsx")
SALIDA_JSON = os.path.join(AQUI, "matriz_litologica_OE4.json")
SALIDA_XLSX = os.path.join(AQUI, "Matriz_litologica_OE4.xlsx")

SGC = "https://srvags.sgc.gov.co/arcgis/rest/services"
CAPAS = {
    "mgc_uc": (SGC + "/Mapa_Geologico_Colombia/Mapa_Geologico_Colombia_V2023/MapServer/733", "MGC 2023 · unidades (1:1.000.000)"),
    "mgc_fallas": (SGC + "/Mapa_Geologico_Colombia/Mapa_Geologico_Colombia_V2023/MapServer/704", "MGC 2023 · fallas"),
    "atlas_uc": (SGC + "/Atlas_Geologico_Colombiano/Atlas_Geologico_Colombia/MapServer/13", "Atlas 2007 · unidades (1:500.000)"),
    "atlas_fallas": (SGC + "/Atlas_Geologico_Colombiano/Atlas_Geologico_Colombia/MapServer/11", "Atlas 2007 · fallas"),
    "tol_uc": (SGC + "/Geologia/Geologia_Por_Departamentos/MapServer/188", "Geología del Tolima · unidades"),
    "tol_fallas": (SGC + "/Geologia/Geologia_Por_Departamentos/MapServer/187", "Geología del Tolima · fallas"),
    "qui_uc": (SGC + "/Geomecanica_Por_Regiones/Geomecanica_del_Departamento_del_Quindio/MapServer/19", "Geología del Quindío · unidades"),
}
A_3116 = pyproj.Transformer.from_crs(4326, 3116, always_xy=True).transform

# ------------------------------------------------------------------ hipótesis [H]
ANCHO_FALLA_M = (60, 350, 590)   # mín / típ / máx de la zona de influencia de una falla: rango y media de La Línea [CP de F, Dávila 2015]
GAMMA_ROCA = 0.027               # MN/m³, peso unitario medio del macizo para p0 = γ·H [H]

# Litologías nuevas (tramos que no coinciden con la Act 3 del OE 3). σci: tabla 1 de Hoek y Marinos (2000) según la
# resistencia de campo de la roca [F] reducida por meteorización/foliación [H]; mi: tabla 2 [F]; GSI: [H].
LITO_NUEVA = {
    "gneis_anfibolita": dict(nombre="Gneises cuarzofeldespáticos y anfibolitas (T-Mmg)", sci=(50, 100, 180), mi=(21, 27, 33), gsi=(35, 50, 65),
                             nota_sci="Gneis y anfibolita: R5-R6 (100 a > 250 MPa) en roca sana; reducido por foliación y fracturamiento [H].",
                             nota_mi="Gneis 28 ± 5; anfibolitas 26 ± 6 (tabla 2).", nota_gsi="Foliado, fracturado; superficies regulares [H]."),
    "esquistos": dict(nombre="Esquistos grafíticos, cuarzomoscovíticos y cloríticos; filitas (J3K1?-Mbg)", sci=(15, 40, 80), mi=(7, 12, 15), gsi=(20, 35, 55),
                      nota_sci="Esquisto R3-R4 (25-100 MPa); el grafito y la sericita lo debilitan: mínimo 15 MPa [H].",
                      nota_mi="Esquistos 12 ± 3; filitas 7 ± 3 (tabla 2); el mínimo toma la filita [H].",
                      nota_gsi="Foliado a laminado/cizallado (tabla 3/4 de Hoek y Marinos); análogo: RMR 50-85 del Complejo Cajamarca en La Línea (lectura cualitativa, Hoek no recomienda convertir RMR a GSI) [H]."),
    "porfido": dict(nombre="Pórfidos dioríticos, granodioríticos y tonalíticos (n4n6-Hi)", sci=(60, 120, 200), mi=(15, 20, 25), gsi=(40, 55, 70),
                    nota_sci="Pórfido: roca ígnea hipoabisal R5 (100-250 MPa) [H].",
                    nota_mi="Pórfidos (20 ± 5), tabla 2.", nota_gsi="Masivo a fracturado; en La Línea el pórfido andesítico dio RMR mayormente 50-80 (lectura de figura) [H]."),
}
ZONA_FALLA = dict(nombre="Zona de falla (roca cataclasada, brecha y salbanda)", sci_factor=0.25, gsi=(10, 15, 25),
                  nota="Zona de falla: σci del tramo × 0,25 y GSI 10-25 [H]; en La Línea las fallas dieron RMR ≈ 20-40 (lectura de figura 4, Castro y Pérez 2013).")

# Tramo MGC → juego de parámetros. Act 3: R1 PCAn, R2 Jcdi, R3 Kqs, R4 Kqv (se leen del Excel).
ASIGNACION = {
    "P-Pi": ["R1", "R2"],          # metagranito/neis pérmico con el Batolito de Ibagué (Act 1-3 del OE 3)
    "J-Pi": ["R2"],
    "T-Mmg": ["gneis_anfibolita"],
    "J3K1?-Mbg": ["esquistos"],
    "n4n6-Hi": ["porfido"],
    "b5k4-VCm": ["R3", "R4"],      # Complejo Quebradagrande: sedimentario y volcánico (Act 3 del OE 3)
}

# Análogo La Línea — Castro y Pérez (2013), tabla 2 (abscisas del túnel piloto, desde el portal Quindío) [F]
TPL_SECTORES = [
    ("A", "Suelo residual", "Suelo residual", 11, 50), ("B", "Complejo Quebradagrande · miembro volcánico", "Diabasas", 50, 1235),
    ("C", "Falla Alaska", "Metadiabasas y lutitas", 1235, 1300), ("D", "Quebradagrande · volcánico sedimentario", "Diabasas, chert, areniscas, lutitas", 1300, 1945),
    ("E", "Falla El Viento", "Metapelitas oscuras grafitosas, cuarcitas y metadiabasas", 1945, 1990),
    ("F", "Quebradagrande · volcánico sedimentario", "Diabasas, chert, areniscas, lutitas", 1990, 2600),
    ("G", "Quebradagrande · sedimentario volcánico", "Areniscas, areniscas conglomeráticas, diabasas, lutitas", 2600, 2935),
    ("H", "Quebradagrande · metasedimentario", "Metalutitas, metalodolitas y metalimolitas", 2935, 2970),
    ("I", "Falla La Vaca", "Lutitas, lodolitas y areniscas finas", 2970, 3160),
    ("J", "Quebradagrande · metasedimentario", "Metalutitas, metalodolitas y metalimolitas", 3160, 3850),
    ("K", "Miembro tectonizado · milonítico", "Esquistos verdes cuarzosericíticos, cuarcitas", 3850, 4335),
    ("L", "Falla La Soledad", "—", 4335, 4870), ("M", "Gabros de La Línea", "Gabros, anfibolitas", 4870, 6080),
    ("N", "Complejo Cajamarca", "Esquistos verdes, esquistos grises", 6080, 7750), ("O", "Falla La Cristalina", "—", 7750, 7830),
    ("P", "Contacto pórfido andesítico / Cajamarca", "Pórfido con xenolitos de esquisto", 7830, 8105),
    ("Q", "Pórfido andesítico", "Pórfido andesítico", 8105, 8542),
]
# Dávila (2015): anchos de falla encontrados en el túnel II Centenario, 8.562 m [F]
DAVILA_FALLAS = [("La Gata", 522), ("Alaska", 185), ("El Viento", 70), ("La Vaca", 436), ("Campanario", 515),
                 ("La Soledad", 590), ("Los Chorros", 440), ("La Cristalina", 60)]
DAVILA_LONGITUD = 8562
# RMR leído de la figura 4 de Castro y Pérez (2013) sobre la imagen renderizada (rangos aproximados) [F, lectura de figura]
TPL_RMR = [("Complejo Quebradagrande (K0-K4+3)", "25-80; mayoría 35-65"), ("Falla La Soledad (K4+3-K4+9)", "20-40"),
           ("Complejo Cajamarca y gabros (K4+9-K8)", "40-85; mayoría 60-80"), ("Fallas señaladas (Alaska, El Viento, La Vaca, La Cristalina)", "≈ 20-35")]

FUENTES = [
    ("MGC23", "Servicio Geológico Colombiano (2023). Mapa Geológico de Colombia, escala 1:1.000.000. Servicio ArcGIS REST, capas 733 (unidades) y 704 (fallas).",
     CAPAS["mgc_uc"][0].rsplit("/", 1)[0]),
    ("ATLAS07", "Gómez Tapias, J. et al. (2007). Atlas Geológico de Colombia, 1:500.000. SGC (INGEOMINAS). Servicio ArcGIS REST, capas 13 y 11.",
     CAPAS["atlas_uc"][0].rsplit("/", 1)[0]),
    ("TOL", "SGC. Geología del departamento del Tolima (servicio «Geología por departamentos», capas 188 y 187). Códigos sin leyenda en el servicio [DA].",
     CAPAS["tol_uc"][0].rsplit("/", 1)[0]),
    ("QUI", "SGC. Geología del departamento del Quindío (servicio «Geomecánica del Quindío», capa 19). Códigos sin leyenda en el servicio [DA].",
     CAPAS["qui_uc"][0].rsplit("/", 1)[0]),
    ("HM00", "Hoek, E. y Marinos, P. (2000). Predicting tunnel squeezing problems in weak heterogeneous rock masses. Tunnels and Tunnelling International. Tablas 1 y 2, ec. 1-2.",
     "https://static.rocscience.cloud/assets/resources/learning/hoek/Predicting-Tunnel-Squeezing-Problems-in-Weak-Heterogeneous-Rock-Masses-2000.pdf"),
    ("CP13", "Castro Caicedo, Á. J. y Pérez Pérez, D. M. (2013). Correlaciones entre las clasificaciones geomecánicas Q y RMR en el túnel exploratorio de «La Línea». Boletín de Ciencias de la Tierra 34, 42-50. Tabla 2 y figura 4.",
     "https://www.redalyc.org/pdf/1695/169530075005.pdf"),
    ("DAV15", "Dávila, H. (2015). Túnel II Centenario, terrenos encontrados y su análisis de comportamiento en el corto y largo plazo. 5ª Jornadas AATES; Revista Vial, 25 nov 2015. Tabla de fallas.",
     "https://revistavial.com/tunel-ii-centenario-terrenos-encontrados-y-su-analisis-de-comportamiento-en-el-corto-y-largo-plazo/"),
    ("ACT3", "Semillero GEOPAV (2026). Cuadro de parámetros adoptados, OE 3 Act 3 (juegos R1-R4, hoja «Roca»).",
     "OE3_Geotecnia/Act3_Parametros/Cuadro_parametros_adoptados_OE3.xlsx"),
    ("OE1", "Semillero GEOPAV (2026). Trazado y perfil del túnel, OE 1 (trazado_tunel.geojson, perfil_continuo.csv).", "OE1_Topografia/"),
    ("OSM", "Vía actual Ruta 40 y Túnel de La Línea según OpenStreetMap (© colaboradores de OSM, ODbL), extraída por via_actual_ruta40.py.",
     "OE2_Plataforma/Act2_Visor/via_actual_ruta40.geojson"),
]


# ------------------------------------------------------------------ utilidades
def eje():
    g = json.load(open(TRAZADO, encoding="utf-8"))["features"][0]["geometry"]
    return g["coordinates"], transform(A_3116, LineString(g["coordinates"]))


def consultar(clave, coords, usar_cache):
    ruta = os.path.join(DATOS, f"sgc_{clave}.geojson")
    if usar_cache and os.path.exists(ruta):
        return json.load(open(ruta, encoding="utf-8"))
    url = CAPAS[clave][0]
    p = {"geometry": json.dumps({"paths": [coords], "spatialReference": {"wkid": 4326}}), "geometryType": "esriGeometryPolyline",
         "inSR": 4326, "outSR": 4326, "spatialRel": "esriSpatialRelIntersects", "outFields": "*", "returnGeometry": "true", "f": "geojson"}
    with urllib.request.urlopen(url + "/query", data=urllib.parse.urlencode(p).encode(), timeout=180) as r:
        g = json.load(r)
    if "error" in g:
        raise SystemExit(f"DETENIDO: el SGC respondió error en {clave}: {g['error']}")
    g["_consulta"] = dict(url=url, fecha=datetime.date.today().isoformat())
    json.dump(g, open(ruta, "w", encoding="utf-8"), ensure_ascii=False)
    return g


def cortes(g, L, poligono):
    """Lista de (pk_ini, pk_fin, props) para polígonos o (pk, None, props, angulo) para líneas."""
    out = []
    for f in g.get("features", []):
        geo = transform(A_3116, shape(f["geometry"]))
        inter = geo.intersection(L)
        if inter.is_empty:
            continue
        partes = list(getattr(inter, "geoms", [inter]))
        props = {k: v for k, v in f["properties"].items() if not k.upper().startswith(("SHAPE", "OBJECTID"))}
        for p in partes:
            if poligono and p.geom_type == "LineString":
                a, b = L.project(Point(p.coords[0])), L.project(Point(p.coords[-1]))
                out.append((min(a, b), max(a, b), props))
            elif not poligono and p.geom_type == "Point":
                pk = L.project(p)
                # ángulo de cruce: dirección local de la falla frente al eje
                seg = geo if geo.geom_type == "LineString" else min(geo.geoms, key=lambda s: s.distance(p))
                d = seg.project(p)
                a, b = seg.interpolate(max(d - 50, 0)), seg.interpolate(min(d + 50, seg.length))
                e1, e2 = L.interpolate(max(pk - 50, 0)), L.interpolate(min(pk + 50, L.length))
                ang_f = math.degrees(math.atan2(b.y - a.y, b.x - a.x))
                ang_e = math.degrees(math.atan2(e2.y - e1.y, e2.x - e1.x))
                ang = abs((ang_f - ang_e + 90) % 180 - 90)
                out.append((pk, None, props, ang))
    return out


def fusionar(segs, clave_props):
    """Une segmentos contiguos con la misma unidad."""
    segs = sorted(segs, key=lambda s: s[0])
    res = []
    for a, b, p in segs:
        k = clave_props(p)
        if res and res[-1][2] == k and a - res[-1][1] < 5:
            res[-1][1] = max(res[-1][1], b)
        else:
            res.append([a, b, k, p])
    return res


def leer_act3():
    wb = load_workbook(ACT3, data_only=True)
    ws = wb["Roca"]
    juegos = {}
    for r in range(6, 18):
        j, caso = ws.cell(r, 1).value, ws.cell(r, 4).value
        if not j:
            continue
        d = juegos.setdefault(j, dict(nombre=ws.cell(r, 3).value, sci=[None] * 3, mi=[None] * 3, gsi=[None] * 3))
        i = ("Mínimo", "Típico", "Máximo").index(caso)
        d["sci"][i], d["mi"][i], d["gsi"][i] = (float(ws.cell(r, c).value) for c in (5, 6, 7))
    if set(juegos) != {"R1", "R2", "R3", "R4"}:
        raise SystemExit(f"DETENIDO: la hoja Roca de la Act 3 no trae R1-R4: {sorted(juegos)}")
    return juegos


def cobertura(pk0, pk1, perfil):
    vals = [c for pk, c in perfil if pk0 <= pk <= pk1]
    if not vals:
        vals = [min(perfil, key=lambda x: abs(x[0] - (pk0 + pk1) / 2))[1]]
    return min(vals), sum(vals) / len(vals), max(vals)


def main():
    usar_cache = "--cache" in sys.argv
    coords, L = eje()
    Ltot = L.length
    perfil = [(float(r["PK_m"]), float(r["Cobertura_m"])) for r in csv.DictReader(open(PERFIL, encoding="utf-8-sig"))]
    crudo = {k: consultar(k, coords, usar_cache) for k in CAPAS}

    # ---------------- sectorización principal (MGC 2023)
    mgc = fusionar(cortes(crudo["mgc_uc"], L, True), lambda p: p["SimboloUC"])
    if abs(mgc[0][0]) > 5 or abs(mgc[-1][1] - Ltot) > 5:
        raise SystemExit("DETENIDO: la sectorización del MGC no cubre el eje completo.")
    for x, y in zip(mgc, mgc[1:]):
        if abs(y[0] - x[1]) > 5:
            raise SystemExit(f"DETENIDO: hueco en la sectorización entre {x[1]:.0f} y {y[0]:.0f} m.")
    act3 = leer_act3()
    tramos = []
    for i, (a, b, simb, p) in enumerate(mgc, start=1):
        if simb not in ASIGNACION:
            raise SystemExit(f"DETENIDO: unidad {simb} sin asignación litológica.")
        cmin, cmed, cmax = cobertura(a, b, perfil)
        juegos = []
        for jg in ASIGNACION[simb]:
            if jg in act3:
                d = act3[jg]
                juegos.append(dict(id=jg, nombre=d["nombre"], sci=d["sci"], mi=d["mi"], gsi=d["gsi"], origen="OE 3 Act 3 (hoja Roca)"))
            else:
                d = LITO_NUEVA[jg]
                juegos.append(dict(id=jg, nombre=d["nombre"], sci=list(d["sci"]), mi=list(d["mi"]), gsi=list(d["gsi"]),
                                   origen="HM00 tablas 1-2 [F] + GSI [H]", notas=[d["nota_sci"], d["nota_mi"], d["nota_gsi"]]))
        tramos.append(dict(tramo=f"T{i}", pk_ini=a, pk_fin=b, longitud=b - a, pct=100 * (b - a) / Ltot, simbolo=simb,
                           descripcion=p["Descripcion"], edad=p["Edad"], cob_min=cmin, cob_med=cmed, cob_max=cmax, juegos=juegos))

    # ---------------- contraste con Atlas 2007, Tolima y Quindío
    def contraste(clave, campo):
        return [dict(pk_ini=a, pk_fin=b, codigo=k) for a, b, k, _ in fusionar(cortes(crudo[clave], L, True), lambda p: p.get(campo))]
    atlas = contraste("atlas_uc", "AUCR_SIMBL")
    tol = contraste("tol_uc", "COD")
    qui = contraste("qui_uc", "COD")

    # ---------------- fallas
    fallas = []
    for clave, fuente, nombre_f, tipo_f in (("mgc_fallas", "MGC 2023", "NombreFalla", "Tipo"), ("atlas_fallas", "Atlas 2007", None, None),
                                            ("tol_fallas", "Tolima", "NMG", None)):
        vistos = []
        for pk, _, props, ang in cortes(crudo[clave], L, False):
            if any(abs(pk - v) < 30 for v in vistos):
                continue
            vistos.append(pk)
            nom = (props.get(nombre_f) or "").strip() if nombre_f else ""
            fallas.append(dict(fuente=fuente, pk=pk, nombre=nom or "sin nombre", tipo=(props.get(tipo_f) if tipo_f else "") or "—", angulo=ang))
    fallas.sort(key=lambda f: (f["fuente"], f["pk"]))
    f_mgc = [f for f in fallas if f["fuente"] == "MGC 2023"]
    for f in f_mgc:
        f["coincide_atlas"] = any(abs(f["pk"] - g["pk"]) < 300 for g in fallas if g["fuente"] == "Atlas 2007")
        f["coincide_tolima"] = any(abs(f["pk"] - g["pk"]) < 300 for g in fallas if g["fuente"] == "Tolima")

    # ---------------- análogo La Línea (geometría)
    via = json.load(open(VIA, encoding="utf-8"))
    tl = max((transform(A_3116, shape(f["geometry"])) for f in via["features"] if f["properties"].get("tunel")), key=lambda g: g.length)
    ext = [Point(tl.coords[0]), Point(tl.coords[-1])]
    linea = dict(longitud_osm=tl.length, pk_extremos=[L.project(p) for p in ext], dist_extremos=[L.distance(p) for p in ext],
                 angulo=abs((math.degrees(math.atan2(tl.coords[-1][1] - tl.coords[0][1], tl.coords[-1][0] - tl.coords[0][0]))
                             - math.degrees(math.atan2(L.coords[-1][1] - L.coords[0][1], L.coords[-1][0] - L.coords[0][0])) + 90) % 180 - 90))
    total_fallas = sum(w for _, w in DAVILA_FALLAS)
    frac_linea = total_fallas / DAVILA_LONGITUD
    frac_mgc = [len(f_mgc) * w / Ltot for w in ANCHO_FALLA_M]

    salida = dict(generado=datetime.datetime.now().isoformat(timespec="seconds"), longitud_eje_m=Ltot, tramos=tramos,
                  contraste=dict(atlas=atlas, tolima=tol, quindio=qui), fallas=fallas,
                  zona_falla=dict(ancho_m=ANCHO_FALLA_M, sci_factor=ZONA_FALLA["sci_factor"], gsi=ZONA_FALLA["gsi"], nota=ZONA_FALLA["nota"]),
                  gamma_roca_MN_m3=GAMMA_ROCA,
                  analogo_linea=dict(geometria=linea, sectores_TPL=TPL_SECTORES, fallas_II_centenario=DAVILA_FALLAS, longitud_II_centenario=DAVILA_LONGITUD,
                                     suma_fallas_m=total_fallas, fraccion_en_falla=frac_linea, rmr_figura4=TPL_RMR),
                  fraccion_en_falla_MGC=frac_mgc,
                  consultas={k: v.get("_consulta") for k, v in crudo.items()})
    json.dump(salida, open(SALIDA_JSON, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    escribir_excel(salida)
    print("escrito", SALIDA_JSON, "y", SALIDA_XLSX)
    for t in tramos:
        print(t["tramo"], t["simbolo"], round(t["pk_ini"]), round(t["pk_fin"]), round(t["cob_min"]), round(t["cob_med"]), round(t["cob_max"]), [j["id"] for j in t["juegos"]])
    for f in f_mgc:
        print("falla", round(f["pk"]), f["nombre"], f["tipo"], round(f["angulo"]), f["coincide_atlas"], f["coincide_tolima"])
    print("La Línea", {k: (round(v) if isinstance(v, float) else [round(x) for x in v]) for k, v in linea.items()}, "fracción en falla", round(frac_linea, 3),
          "MGC", [round(x, 3) for x in frac_mgc])


# ------------------------------------------------------------------ Excel
VERDE, AZUL, GRIS, ROJO = "178E2C", "193F77", "475569", "9A3412"
F = "Arial"
_fino = Side(style="thin", color="CBD5E1")
BORDE = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)
NORMAL, NEGRA = Font(name=F, size=10), Font(name=F, size=10, bold=True)
ENTRADA, ENLACE = Font(name=F, size=10, color="0000FF"), Font(name=F, size=10, color="008000")
CAB = Font(name=F, size=10, bold=True, color="FFFFFF")
FILL_CAB, FILL_H, FILL_RES = PatternFill("solid", fgColor=AZUL), PatternFill("solid", fgColor="FFF7D6"), PatternFill("solid", fgColor="E8F3EA")


def _cab(ws, fila, textos, anchos=None):
    for j, t in enumerate(textos, start=1):
        c = ws.cell(fila, j, t)
        c.font, c.fill, c.border = CAB, FILL_CAB, BORDE
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    if anchos:
        for j, w in enumerate(anchos, start=1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = w
    ws.row_dimensions[fila].height = 40


def _c(ws, r, c, v, fuente=NORMAL, fmt=None, h=False, wrap=False, fill=None):
    x = ws.cell(r, c, v)
    x.font, x.border = fuente, BORDE
    x.alignment = Alignment(wrap_text=wrap, vertical="top")
    if fmt:
        x.number_format = fmt
    if h:
        x.fill = FILL_H
    if fill:
        x.fill = fill
    return x


def _tit(ws, t, fila=1):
    ws.cell(fila, 1, t).font = Font(name=F, size=13, bold=True, color=AZUL)


def _nota(ws, r, t, color=GRIS):
    ws.cell(r, 1, t).font = Font(name=F, size=9, italic=True, color=color)


def escribir_excel(S):
    wb = Workbook()
    p = wb.active
    p.title = "Portada"
    p.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", (3, 24, 22, 22, 22, 22, 22, 3)):
        p.column_dimensions[col].width = w
    for nombre, cel in (("logo-unibague.png", "B2"), ("logo-geopav.png", "F2")):
        ruta = os.path.join(RAIZ, "public", nombre)
        if os.path.exists(ruta):
            im = XLImage(ruta)
            k = 70 / im.height
            im.height, im.width = 70, int(im.width * k)
            p.add_image(im, cel)
    p["B8"] = "Matriz litológica del macizo a lo largo del túnel"
    p["B8"].font = Font(name=F, size=18, bold=True, color=VERDE)
    p["B9"] = "Objetivo específico 4 · Actividad 1 — Caracterización litológica y de zonas de falla a partir de la cartografía del SGC"
    p["B9"].font = Font(name=F, size=12, color=AZUL)
    ficha = [
        ("Objetivo", "OE 4: Estructurar la metodología de excavación recomendada, contrastando el método mecanizado con tuneladora frente al método convencional, a partir de la caracterización geológica regional del macizo"),
        ("Actividad", "1. Caracterización litológica y de zonas de falla del macizo a partir de la cartografía del SGC"),
        ("Entregable / formato", "Matriz litológica · Documento Excel"),
        ("Periodo", "Bloque 2 del Plan de Acción (5 oct - 7 nov 2026)"),
        ("Responsable asignado", "Castaño Cifuentes Maicol Stiven"),
        ("Apoyo en la ejecución", "Tamayo Osorio Miguel Ángel, con apoyo de IA (Claude)"),
        ("Semillero", "Semillero de Investigación GEOPAV · Ingeniería Civil · Universidad de Ibagué · Paz y Región 2026B"),
        ("Fecha de elaboración", "29 sep 2026"),
        ("Advertencia", "Caracterización REGIONAL (escalas 1:1.000.000 a 1:500.000). No hay sondeos a la profundidad del túnel: los parámetros de macizo son rangos de fuentes publicadas e hipótesis declaradas. "
                        "Las fallas cartografiadas a esta escala subestiman las que encontrará la excavación (análogo: Túnel de La Línea)."),
    ]
    for i, (k, v) in enumerate(ficha):
        r = 11 + i
        p.cell(r, 2, k).font = NEGRA
        p.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)
        p.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        c = p.cell(r, 3, v)
        c.font = NORMAL if k != "Advertencia" else Font(name=F, size=10, bold=True, color=ROJO)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        p.row_dimensions[r].height = 30 if len(v) < 110 else 58
    r = 11 + len(ficha) + 1
    p.cell(r, 2, "Índice de hojas (tabla de contenido)").font = Font(name=F, size=12, bold=True, color=AZUL)
    r += 1
    for h_, d in [("Matriz", "Tabla 1. Matriz litológica por tramos del eje (MGC 2023) con cobertura y parámetros de macizo"),
                  ("Parámetros", "Tabla 2. Juegos de parámetros de macizo (σci, mi, GSI) y su origen"),
                  ("Fallas", "Tabla 3. Fallas que cortan el eje según tres cartografías · Tabla 4. Zona de influencia adoptada"),
                  ("Contraste", "Tabla 5. Unidades según Atlas 2007, Tolima y Quindío frente al MGC 2023"),
                  ("Análogo La Línea", "Tabla 6. Sectores del túnel piloto · Tabla 7. Fallas del túnel II Centenario · Tabla 8. RMR observado"),
                  ("Marcas y fuentes", "Tabla 9. Marcas F / CP / H / DA · Tabla 10. Fuentes citadas y consultas")]:
        p.cell(r, 2, h_).font = NEGRA
        p.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        p.cell(r, 3, d).font = NORMAL
        r += 1
    r += 1
    p.merge_cells(start_row=r, start_column=2, end_row=r + 3, end_column=7)
    c = p.cell(r, 2, "Índice de tablas: tablas 1 a 10 según el índice de hojas. Índice de figuras: no tiene figuras (el perfil geológico va en el plano de la Act 2). "
                     "Índice de anexos: no tiene anexos; las respuestas crudas del SGC quedan en OE4_Tuneles/Act1_Litologia/Datos/. "
                     "Convención: verde = valor leído de otro producto; fondo amarillo = hipótesis [H]. Abscisa (PK) en metros desde el portal oriental (Ibagué). "
                     "Elaborado con apoyo de IA (Claude) y verificado contra las fuentes. Script: OE4_Tuneles/Act1_Litologia/matriz_litologica_OE4.py")
    c.font = Font(name=F, size=9, italic=True, color=GRIS)
    c.alignment = Alignment(wrap_text=True, vertical="top")

    # ---- Matriz
    m = wb.create_sheet("Matriz")
    m.sheet_view.showGridLines = False
    _tit(m, "Tabla 1. Matriz litológica del eje del túnel por tramos del Mapa Geológico de Colombia 2023")
    _nota(m, 2, "PK en metros desde el portal oriental. Cobertura = cota del terreno − cota de rasante (perfil del OE 1) [CP]. Unidades del MGC 2023 [F]. Parámetros: hoja «Parámetros».")
    _cab(m, 4, ["Tramo", "PK inicio (m)", "PK fin (m)", "Longitud (m)", "% del eje", "Símbolo MGC", "Litología (MGC 2023)", "Edad",
                "Cobertura mín. (m)", "Cobertura media (m)", "Cobertura máx. (m)", "Juegos de parámetros", "Fallas MGC en el tramo o en su borde"],
         [7, 11, 11, 11, 8, 13, 50, 20, 11, 11, 11, 18, 34])
    f_mgc = [f for f in S["fallas"] if f["fuente"] == "MGC 2023"]
    for i, t in enumerate(S["tramos"]):
        r = 5 + i
        ff = [f"{f['nombre']} (PK {f['pk']:,.0f})".replace(",", ".") for f in f_mgc if t["pk_ini"] - 30 <= f["pk"] <= t["pk_fin"] + 30]
        vals = [t["tramo"], t["pk_ini"], t["pk_fin"], f"=C{r}-B{r}", f"=D{r}/{S['longitud_eje_m']:.1f}", t["simbolo"], t["descripcion"], t["edad"],
                t["cob_min"], t["cob_med"], t["cob_max"], " + ".join(j["id"] for j in t["juegos"]), "; ".join(ff) or "—"]
        for j, v in enumerate(vals, start=1):
            fmt = {2: "#,##0", 3: "#,##0", 4: "#,##0", 5: "0.0%", 9: "#,##0", 10: "#,##0", 11: "#,##0"}.get(j)
            _c(m, r, j, v, NEGRA if j == 1 else NORMAL, fmt, wrap=j in (7, 13))
        m.row_dimensions[r].height = 44
    r = 5 + len(S["tramos"])
    _c(m, r, 1, "Total", NEGRA)
    _c(m, r, 4, f"=SUM(D5:D{r - 1})", NEGRA, "#,##0", fill=FILL_RES)
    _c(m, r, 5, f"=SUM(E5:E{r - 1})", NEGRA, "0.0%", fill=FILL_RES)
    _nota(m, r + 2, "Fuente: MGC 2023 (capas 733 y 704), consultado por el script; perfil del OE 1. La cobertura supera los 1.000 m en buena parte de los tramos T4-T6: allí el esfuerzo in situ, no solo la litología, decide el método (Act 2).")
    m.freeze_panes = "B5"

    # ---- Parámetros
    pa = wb.create_sheet("Parámetros")
    pa.sheet_view.showGridLines = False
    _tit(pa, "Tabla 2. Juegos de parámetros del macizo rocoso por litología (mínimo / típico / máximo)")
    _nota(pa, 2, "R1-R4 se leen del Excel de la Act 3 del OE 3 (texto verde). Los demás: σci por resistencia de campo (HM00 tabla 1) [F→H], mi de HM00 tabla 2 [F], GSI [H].")
    _cab(pa, 4, ["Juego", "Litología", "σci mín (MPa)", "σci típ", "σci máx", "mi mín", "mi típ", "mi máx", "GSI mín", "GSI típ", "GSI máx", "Origen", "Notas"],
         [18, 46, 10, 9, 9, 8, 8, 8, 8, 8, 8, 24, 70])
    vistos, r = set(), 5
    for t in S["tramos"]:
        for jg in t["juegos"]:
            if jg["id"] in vistos:
                continue
            vistos.add(jg["id"])
            desde_act3 = jg["origen"].startswith("OE 3")
            vals = [jg["id"], jg["nombre"], *jg["sci"], *jg["mi"], *jg["gsi"], jg["origen"], " ".join(jg.get("notas", [])) or "Ver notas por juego en la Act 3 del OE 3."]
            for j, v in enumerate(vals, start=1):
                _c(pa, r, j, v, ENLACE if desde_act3 and 3 <= j <= 11 else (NEGRA if j == 1 else NORMAL), "0" if 3 <= j <= 11 else None,
                   h=(not desde_act3 and j in (3, 4, 5, 9, 10, 11)), wrap=j in (2, 13))
            pa.row_dimensions[r].height = 54
            r += 1
    _c(pa, r, 1, "Zona de falla", NEGRA)
    _c(pa, r, 2, ZONA_FALLA["nombre"], wrap=True)
    _c(pa, r, 3, f"σci del tramo × {ZONA_FALLA['sci_factor']}", h=True)
    for j, v in zip((9, 10, 11), ZONA_FALLA["gsi"]):
        _c(pa, r, j, v, NORMAL, "0", h=True)
    _c(pa, r, 12, "[H] + análogo CP13")
    _c(pa, r, 13, ZONA_FALLA["nota"], wrap=True)
    pa.row_dimensions[r].height = 44
    _nota(pa, r + 2, f"Peso unitario del macizo para el esfuerzo in situ p0 = γ·H: {S['gamma_roca_MN_m3']} MN/m³ [H]. Estos juegos alimentan el índice de squeezing de la Act 2 (Hoek y Marinos, 2000).")

    # ---- Fallas
    fa = wb.create_sheet("Fallas")
    fa.sheet_view.showGridLines = False
    _tit(fa, "Tabla 3. Fallas cartografiadas que cortan el eje del túnel")
    _nota(fa, 2, "Ángulo de cruce entre la traza de la falla y el eje, medido sobre ±50 m [CP]. «Coincide» = otra cartografía corta el eje a menos de 300 m.")
    _cab(fa, 4, ["Fuente", "PK (m)", "Nombre", "Tipo", "Ángulo de cruce (°)", "Coincide Atlas 2007", "Coincide Tolima"], [12, 11, 30, 34, 13, 13, 13])
    r = 5
    for f in S["fallas"]:
        vals = [f["fuente"], f["pk"], f["nombre"], f["tipo"], f["angulo"],
                ("Sí" if f.get("coincide_atlas") else "No") if f["fuente"] == "MGC 2023" else "",
                ("Sí" if f.get("coincide_tolima") else "No") if f["fuente"] == "MGC 2023" else ""]
        for j, v in enumerate(vals, start=1):
            _c(fa, r, j, v, NEGRA if f["fuente"] == "MGC 2023" and j == 3 else NORMAL, {2: "#,##0", 5: "0"}.get(j), wrap=j == 4)
        r += 1
    r += 1
    fa.cell(r, 1, "Tabla 4. Zona de influencia adoptada por falla y fracción del eje en zona de falla").font = Font(name=F, size=13, bold=True, color=AZUL)
    r += 1
    _cab(fa, r, ["Caso", "Ancho por falla (m)", "Fallas MGC", "Longitud en falla (m)", "% del eje", "Referencia"])
    n = len([f for f in S["fallas"] if f["fuente"] == "MGC 2023"])
    r0 = r + 1
    for i, (caso, w) in enumerate(zip(("Mínimo", "Típico", "Máximo"), S["zona_falla"]["ancho_m"])):
        rr = r0 + i
        _c(fa, rr, 1, caso, NEGRA)
        _c(fa, rr, 2, w, ENTRADA, "#,##0", h=True)
        _c(fa, rr, 3, n, NORMAL)
        _c(fa, rr, 4, f"=B{rr}*C{rr}", NORMAL, "#,##0")
        _c(fa, rr, 5, f"=D{rr}/{S['longitud_eje_m']:.1f}", NORMAL, "0.0%", fill=FILL_RES)
        _c(fa, rr, 6, "Mín., media y máx. de los anchos de falla de La Línea (DAV15) [CP]", wrap=True)
    rr = r0 + 3
    _c(fa, rr, 1, "La Línea (medido)", NEGRA)
    _c(fa, rr, 4, S["analogo_linea"]["suma_fallas_m"], ENLACE, "#,##0")
    _c(fa, rr, 5, S["analogo_linea"]["fraccion_en_falla"], NEGRA, "0.0%", fill=FILL_RES)
    _c(fa, rr, 6, f"8 fallas en {S['analogo_linea']['longitud_II_centenario']:,} m del túnel II Centenario (DAV15) [F→CP]".replace(",", "."), wrap=True)
    _nota(fa, rr + 2, "Lectura: con solo las fallas del mapa 1:1.000.000 el eje tendría del orden de 1-6 % en zona de falla; el túnel de La Línea, en el mismo macizo, encontró un 33 %. "
                      "La Act 2 usa las fallas cartografiadas y declara esta subestimación como sensibilidad.", ROJO)

    # ---- Contraste
    co = wb.create_sheet("Contraste")
    co.sheet_view.showGridLines = False
    _tit(co, "Tabla 5. Unidades cortadas por el eje según otras cartografías del SGC (para contrastar con el MGC 2023)")
    _nota(co, 2, "Tolima y Quindío: el servicio no publica la leyenda de los códigos [DA]. Se listan para mostrar dónde hay mayor detalle y depósitos cuaternarios superficiales, que no alcanzan la rasante salvo cerca de los portales.")
    _cab(co, 4, ["Cartografía", "PK inicio (m)", "PK fin (m)", "Longitud (m)", "Código / símbolo"], [16, 12, 12, 12, 22])
    r = 5
    for nom, lst in (("Atlas 2007", S["contraste"]["atlas"]), ("Tolima", S["contraste"]["tolima"]), ("Quindío", S["contraste"]["quindio"])):
        for x in lst:
            for j, v in enumerate([nom, x["pk_ini"], x["pk_fin"], f"=C{r}-B{r}", x["codigo"]], start=1):
                _c(co, r, j, v, NORMAL, "#,##0" if j in (2, 3, 4) else None)
            r += 1

    # ---- Análogo La Línea
    an = wb.create_sheet("Análogo La Línea")
    an.sheet_view.showGridLines = False
    g = S["analogo_linea"]["geometria"]
    _tit(an, "Tabla 6. Túnel piloto de La Línea: sectores geológicos encontrados (Castro y Pérez, 2013, tabla 2)")
    _nota(an, 2, f"Relación con el eje de la Ferropista [CP]: el Túnel de La Línea (OSM, {g['longitud_osm']:,.0f} m) tiene un extremo a {g['dist_extremos'][0]:,.0f} m en planta del eje (PK {g['pk_extremos'][0]:,.0f}) "
                 f"y el otro a {g['dist_extremos'][1]:,.0f} m (PK {g['pk_extremos'][1]:,.0f}); ángulo entre ejes {g['angulo']:.0f}°. Está unos 1.000 m por encima de la rasante.".replace(",", "."))
    _cab(an, 4, ["Sector", "Unidad", "Litologías", "Abscisa inicio (m)", "Abscisa fin (m)", "Longitud (m)"], [8, 40, 44, 14, 14, 12])
    r = 5
    for s in S["analogo_linea"]["sectores_TPL"]:
        for j, v in enumerate([s[0], s[1], s[2], s[3], s[4], f"=E{r}-D{r}"], start=1):
            _c(an, r, j, v, NEGRA if j == 1 else NORMAL, "#,##0" if j >= 4 else None, wrap=j in (2, 3))
        r += 1
    _nota(an, r, "Fuente: CP13 tabla 2 (abscisas desde el portal Quindío, túnel piloto de 8.554 m) [F].")
    r += 2
    an.cell(r, 1, "Tabla 7. Zonas de falla encontradas en el túnel II Centenario (Dávila, 2015)").font = Font(name=F, size=13, bold=True, color=AZUL)
    r += 1
    _cab(an, r, ["Falla", "Longitud atravesada (m)"])
    r0 = r + 1
    for i, (nom, w) in enumerate(S["analogo_linea"]["fallas_II_centenario"]):
        _c(an, r0 + i, 1, nom, NEGRA)
        _c(an, r0 + i, 2, w, NORMAL, "#,##0")
    rt = r0 + len(S["analogo_linea"]["fallas_II_centenario"])
    _c(an, rt, 1, "Total", NEGRA)
    _c(an, rt, 2, f"=SUM(B{r0}:B{rt - 1})", NEGRA, "#,##0", fill=FILL_RES)
    _c(an, rt + 1, 1, "Fracción del túnel", NEGRA)
    _c(an, rt + 1, 2, f"=B{rt}/{S['analogo_linea']['longitud_II_centenario']}", NEGRA, "0.0%", fill=FILL_RES)
    _nota(an, rt + 2, "Fuente: DAV15 [F]; total y fracción [CP]. Longitud del túnel II Centenario: 8.562 m (DAV15).")
    r = rt + 4
    an.cell(r, 1, "Tabla 8. RMR observado en el túnel piloto (lectura de la figura 4 de Castro y Pérez, 2013)").font = Font(name=F, size=13, bold=True, color=AZUL)
    r += 1
    _cab(an, r, ["Sector", "RMR (rango leído)"])
    for i, (a, b) in enumerate(S["analogo_linea"]["rmr_figura4"]):
        _c(an, r + 1 + i, 1, a, NORMAL, wrap=True)
        _c(an, r + 1 + i, 2, b)
    _nota(an, r + 2 + len(S["analogo_linea"]["rmr_figura4"]), "Lectura visual sobre la imagen renderizada de la figura [F, lectura aproximada]. Hoek (cap. 11) desaconseja convertir RMR a GSI: se usa solo como contraste cualitativo.")

    # ---- Marcas y fuentes
    mf = wb.create_sheet("Marcas y fuentes")
    mf.sheet_view.showGridLines = False
    _tit(mf, "Tabla 9. Marcas de origen")
    _cab(mf, 3, ["Marca", "Significado"], [12, 110])
    for i, (a, b) in enumerate((("F", "Dato de fuente publicada o de servicio oficial, leído en el original"),
                                ("CP", "Cálculo propio, reproducible con el script citado"),
                                ("H", "Hipótesis del semillero, declarada; fondo amarillo"),
                                ("DA", "Dato abierto identificado y pendiente de obtener"))):
        _c(mf, 4 + i, 1, a, NEGRA)
        _c(mf, 4 + i, 2, b)
    mf["A10"] = "Tabla 10. Fuentes citadas y consultas a servicios"
    mf["A10"].font = Font(name=F, size=13, bold=True, color=AZUL)
    _cab(mf, 11, ["Clave", "Referencia", "Enlace / ruta"])
    mf.row_dimensions[11].height = 20
    mf.column_dimensions["C"].width = 80
    for i, (a, b, c_) in enumerate(FUENTES):
        _c(mf, 12 + i, 1, a, NEGRA)
        _c(mf, 12 + i, 2, b, wrap=True)
        _c(mf, 12 + i, 3, c_, wrap=True)
        mf.row_dimensions[12 + i].height = 34
    r = 12 + len(FUENTES) + 1
    for k, v in S["consultas"].items():
        if v:
            _nota(mf, r, f"Consulta {k}: {v['url']} · {v['fecha']}")
            r += 1

    for hoja in wb.worksheets:
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.fitToWidth, hoja.page_setup.fitToHeight = 1, 0
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(SALIDA_XLSX)


if __name__ == "__main__":
    main()
