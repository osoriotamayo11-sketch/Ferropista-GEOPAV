# -*- coding: utf-8 -*-
"""
OE 5 - Actividad 1
Extraccion de las series de volumenes de transito de INVIAS para el corredor
Ibague - Calarca (Ruta Nacional 40, paso del Alto de La Linea).

Semillero de Investigacion GEOPAV - Universidad de Ibague - Paz y Region 2026B

QUE HACE
  Lee los dos libros publicados por INVIAS y extrae unicamente las dos estaciones
  de aforo que cubren el corredor de analisis, mas el aforo de peaje de 2019.
  Escribe dos CSV con separador ';' y formato numerico es-CO.

FUENTES  [F]
  Serie historica de transito (TPD) 1997-2018  - INVIAS, idFile=30450
  Serie historica de transito (TPD) 2019       - INVIAS, idFile=1678
  https://www.invias.gov.co/publicaciones/4154/documentos-tecnicos/
  Descargados el 8 sep 2026. Los enlaces indexados antiguos
  (/archivo-y-documentos/documentos-tecnicos/volumenes-de-transito) responden 404.

ESTACIONES DEL CORREDOR
  243  TOLIMA   CAJAMARCA-IBAGUE                 L = 29 km   dentro del corredor
  244  QUINDIO  CALARCA - CAJAMARCA   via 4003   L = 45 km   dentro - el paso de La Linea
  245  QUINDIO  ARMENIA - CALARCA     via 4003   L =  5 km   FUERA del corredor entre portales

  El corredor de analisis del OE 5 va entre los portales del tunel (OE 1), y el portal
  occidental esta en CALARCA, no en Armenia. Por eso el tramo Armenia - Calarca queda
  fuera: L = 29 + 45 = 74 km  [F, INVIAS].
  La estacion 245 se extrae de todos modos y se marca en_corredor_portales = No, para
  que la exclusion sea una decision visible y no una omision silenciosa. Si alguna vez
  se analiza el corredor completo Ibague - Armenia, son 79 km.
  Se descarta tambien la estacion 1217 (Calarca - Av. Centenario, 4 km): via urbana.

  El codigo de via se copia tal como lo trae el libro. La estacion 243 lo trae VACIO:
  se deja vacio y no se completa por deduccion.

DEFECTO DE ORIGEN CORREGIDO AQUI
  La fila de composicion vehicular ("autos-buses-camiones", p. ej. "35-10-55") fue
  interpretada por Excel como fecha en varios anios: "22-09-69" quedo guardado como
  1969-09-22 y otros como numero de serie. Se revierte leyendo dia-mes-anio y se
  valida que los tres porcentajes sumen 100 +/- 1. Lo que no valida se deja vacio.

ADVERTENCIA - SON DOS PEAJES DISTINTOS, NO COMPARAR SIN HOMOLOGAR  [F]
  Peaje CAJAMARCA  - PR 24+020 de la carretera Armenia - Ibague, administrado por
                     INVIAS. Ficha oficial: hermes.invias.gov.co/recursos/
                     Fichas%20Peajes/3985.pdf
  Peaje COCORA     - K 13+750 de la Segunda Calzada Ibague - Cajamarca, corregimiento
                     Coello Cocora, jurisdiccion de Ibague. Administrado por la
                     concesion APP GICA (ANI). Fuente: appgica.com.co/peajes
  Estan separados unos 10 km y Cocora, mas abajo, capta trafico local que no sube al
  paso. La serie de la ANI que usa la memoria del OE 5 es la de COCORA; la de INVIAS
  que extrae este script es la de CAJAMARCA. Cualquier comparacion entre las dos
  series exige homologar antes el punto de medicion y las categorias tarifarias.
"""

import csv
import datetime
import os
import openpyxl

# El script vive en Act1_Aforos/ y lee y escribe en esa misma carpeta.
BASE = os.path.dirname(os.path.abspath(__file__))
DATOS = BASE

LIBRO_HIST = os.path.join(DATOS, "TPD_INVIAS_1997_2018.xlsx")
LIBRO_2019 = os.path.join(DATOS, "TPD_INVIAS_2019_peajes.xlsx")

SALIDA_HIST = os.path.join(DATOS, "tpd_historico_corredor.csv")
SALIDA_2019 = os.path.join(DATOS, "tpd_2019_peaje_cajamarca.csv")

# hoja, fila (0-based), desplazamiento de la columna SECTOR, metadatos declarados
ESTACIONES = [
    {"hoja": "TOLIMA",  "fila": 35, "off": 2, "estacion": "243",
     "sector": "CAJAMARCA-IBAGUE",    "long_km": 29, "en_corredor": True},
    {"hoja": "QUINDIO", "fila": 5,  "off": 1, "estacion": "244",
     "sector": "CALARCA - CAJAMARCA", "long_km": 45, "en_corredor": True},
    {"hoja": "QUINDIO", "fila": 7,  "off": 1, "estacion": "245",
     "sector": "ARMENIA - CALARCA",   "long_km": 5,  "en_corredor": False},
]

# Categorias tarifarias del aforo de peaje. III, IV y V son carga pesada.
CAT_LIVIANO = ["I"]
CAT_INTERMEDIO = ["II"]
CAT_PESADO = ["III", "IV", "V", "VI", "VII"]


def es_co(valor, decimales=1):
    """Formato numerico colombiano: coma decimal, punto de miles."""
    if valor is None:
        return ""
    s = f"{valor:,.{decimales}f}"
    return s.replace(",", "@").replace(".", ",").replace("@", ".")


def recuperar_composicion(bruto):
    """
    Devuelve (pct_autos, pct_buses, pct_camiones) o (None, None, None).

    INVIAS publica la composicion como texto "AA-BB-CC". Excel convirtio parte de
    esos valores a fecha. Se revierte: dia = autos, mes = buses, anio de dos
    cifras = camiones. Solo se acepta si los tres suman 100 +/- 1.
    """
    if bruto is None or bruto == "":
        return (None, None, None)

    trio = None

    if isinstance(bruto, datetime.datetime):
        trio = (bruto.day, bruto.month, bruto.year % 100)
    elif isinstance(bruto, (int, float)) and 1 < bruto < 60000:
        # numero de serie de Excel (base 1899-12-30)
        f = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=int(bruto))
        trio = (f.day, f.month, f.year % 100)
    else:
        partes = [p for p in str(bruto).replace("--", "").split("-") if p.strip().isdigit()]
        if len(partes) == 3:
            trio = tuple(int(p) for p in partes)

    if trio is None:
        return (None, None, None)
    if abs(sum(trio) - 100) > 1:
        return (None, None, None)
    return trio


def leer_serie_historica():
    wb = openpyxl.load_workbook(LIBRO_HIST, read_only=True, data_only=True)
    filas = []
    descartes = 0

    for est in ESTACIONES:
        ws = wb[est["hoja"]]
        datos = [list(r) for r in ws.iter_rows(values_only=True)]
        cabecera = datos[2]
        anios = [(i, int(str(v)[:4])) for i, v in enumerate(cabecera)
                 if v is not None and str(v)[:4].isdigit()]

        fila_tpd = datos[est["fila"]]
        fila_comp = datos[est["fila"] + 1]

        # el codigo de via se COPIA del libro; si viene vacio, queda vacio
        via_libro = fila_tpd[est["off"] + 1]
        via = str(via_libro).strip() if via_libro not in (None, "") else ""

        # el sector declarado en el libro debe coincidir con el esperado
        sector_libro = str(fila_tpd[est["off"]] or "").strip().upper()
        esperado = est["sector"].replace(" ", "").upper()
        if sector_libro.replace(" ", "")[:12] != esperado[:12]:
            raise SystemExit(
                f"PARADA: la fila {est['fila']} de {est['hoja']} dice "
                f"{sector_libro!r}, no {est['sector']!r}. El libro cambio de "
                f"estructura; revisar antes de continuar."
            )

        for col, anio in anios:
            tpd = fila_tpd[col]
            if not isinstance(tpd, (int, float)) or tpd <= 0:
                continue
            a, b, c = recuperar_composicion(fila_comp[col])
            if a is None and fila_comp[col] not in (None, "", "--"):
                descartes += 1
            filas.append({
                "estacion": est["estacion"],
                "departamento": est["hoja"],
                "sector": est["sector"],
                "codigo_via": via,
                "longitud_km": est["long_km"],
                "en_corredor_portales": "Si" if est["en_corredor"] else "No",
                "anio": anio,
                "tpd": int(tpd),
                "pct_autos": a,
                "pct_buses": b,
                "pct_camiones": c,
                "tpd_camiones": round(tpd * c / 100, 1) if c is not None else None,
            })

    filas.sort(key=lambda f: (f["estacion"], f["anio"]))
    return filas, descartes


def leer_peaje_2019():
    wb = openpyxl.load_workbook(LIBRO_2019, read_only=True, data_only=True)
    ws = wb["Info Peajes"]
    cabecera = None
    meses = []

    for r in ws.iter_rows(values_only=True):
        if cabecera is None:
            cabecera = [str(x).strip().upper() if x is not None else "" for x in r]
            continue
        if str(r[4] or "").strip().upper() != "CAJAMARCA":
            continue
        if str(r[2] or "").strip().upper() != "TOLIMA":
            continue
        cats = {}
        for i, nombre in enumerate(cabecera):
            if nombre.startswith("CATEGORIA"):
                v = r[i]
                cats[nombre.replace("CATEGORIA", "").strip()] = v if isinstance(v, (int, float)) else 0
        desde = r[5]
        meses.append({
            "mes": desde.strftime("%Y-%m") if isinstance(desde, datetime.datetime) else str(desde)[:7],
            "dias": (r[6] - desde).days + 1 if isinstance(desde, datetime.datetime) else None,
            "cats": cats,
        })

    if len(meses) != 12:
        raise SystemExit(f"PARADA: se esperaban 12 meses del peaje Cajamarca y hay {len(meses)}.")

    filas = []
    for m in meses:
        liv = sum(m["cats"].get(c, 0) for c in CAT_LIVIANO)
        inter = sum(m["cats"].get(c, 0) for c in CAT_INTERMEDIO)
        pes = sum(m["cats"].get(c, 0) for c in CAT_PESADO)
        total = liv + inter + pes
        filas.append({
            "mes": m["mes"], "dias": m["dias"], "total": total,
            "livianos": liv, "intermedios": inter, "pesados": pes,
            "tpd": round(total / m["dias"], 1) if m["dias"] else None,
            "tpd_pes": round(pes / m["dias"], 1) if m["dias"] else None,
        })
    return filas


def escribir(ruta, campos, filas, formateo):
    with open(ruta, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(campos)
        for f in filas:
            w.writerow([formateo(f, c) for c in campos])


def main():
    hist, descartes = leer_serie_historica()
    peaje = leer_peaje_2019()

    campos_h = ["estacion", "departamento", "sector", "codigo_via", "longitud_km",
                "en_corredor_portales", "anio", "tpd", "pct_autos", "pct_buses",
                "pct_camiones", "tpd_camiones"]

    def fmt_h(f, c):
        v = f[c]
        if v is None:
            return ""
        if c == "tpd_camiones":
            return es_co(v, 1)
        return str(v)

    campos_p = ["mes", "dias", "total", "livianos", "intermedios", "pesados", "tpd", "tpd_pes"]

    def fmt_p(f, c):
        v = f[c]
        if v is None:
            return ""
        if c in ("tpd", "tpd_pes"):
            return es_co(v, 1)
        return str(v)

    escribir(SALIDA_HIST, campos_h, hist, fmt_h)
    escribir(SALIDA_2019, campos_p, peaje, fmt_p)

    # ---- resumen en consola, para pegar en la memoria ----
    print(f"Serie historica: {len(hist)} registros de {len(ESTACIONES)} estaciones.")
    if descartes:
        print(f"  {descartes} valores de composicion no validaron la suma 100 y quedaron vacios.")

    print("\nPeriodo de los fallecidos de la ANSV (2015-2019):")
    for est in ESTACIONES:
        sub = [f for f in hist if f["estacion"] == est["estacion"] and 2015 <= f["anio"] <= 2019]
        if not sub:
            continue
        media = sum(f["tpd"] for f in sub) / len(sub)
        anios = ", ".join(f"{f['anio']}={f['tpd']}" for f in sub)
        marca = "" if est["en_corredor"] else "   [FUERA del corredor entre portales]"
        print(f"  Est.{est['estacion']} {est['sector']} (L={est['long_km']} km){marca}")
        print(f"    {anios}")
        print(f"    TPD medio -> {es_co(media, 0)} veh/dia")

    L = sum(e["long_km"] for e in ESTACIONES if e["en_corredor"])
    print(f"\nLongitud del corredor entre portales (est. 243 + 244): {L} km  [F, INVIAS]")

    tot = sum(f["total"] for f in peaje)
    pes = sum(f["pesados"] for f in peaje)
    dias = sum(f["dias"] for f in peaje)
    print(f"\nPeaje Cajamarca 2019 ({dias} dias): {tot:,} veh".replace(",", "."))
    print(f"  TPD {es_co(tot/dias, 0)} veh/dia | pesados {es_co(pes/dias, 0)} veh/dia "
          f"| participacion {es_co(pes/tot*100, 1)} %")
    print(f"\nEscritos:\n  {SALIDA_HIST}\n  {SALIDA_2019}")


if __name__ == "__main__":
    main()
