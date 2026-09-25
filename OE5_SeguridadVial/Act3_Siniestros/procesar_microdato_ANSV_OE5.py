# -*- coding: utf-8 -*-
"""
OE 5 - Actividad 3 (fila 33) - Microdato georreferenciado de la ANSV, 2021 - mar 2026
Semillero de Investigacion GEOPAV - Universidad de Ibague - Paz y Region 2026B

FUENTE [F]
  ANSV, Observatorio Nacional de Seguridad Vial. Respuesta al radicado 20266600072342,
  oficio 20265000140371 del 22 sep 2026 (Resp_20265000140371.pdf) y su anexo
  Sint_Via_LaLinea_CodTramo_4003_Vigencia_2021_Mar2026.xlsx. Fuente primaria de la ANSV:
  INMLCF (SIRDEC / SICLICO). Solo incluye los registros que la ANSV pudo georreferenciar.

QUE HACE (lee, no escribe nada a mano)
  1. Lee la hoja de registros (una fila por VICTIMA) y comprueba que sus conteos
     reproducen las tablas dinamicas que la propia ANSV incluyo en el anexo. Si no, aborta.
  2. Agrupa victimas en HECHOS. [H] Un hecho = misma fecha + mismo municipio + mismo PR
     y distancia. El anexo no trae identificador de siniestro.
  3. Abscisa continua de la Ruta 4003 [CP]: km = PR + distancia/1000. El anexo declara
     que el sector "Armenia - La Linea" termina en PR 50+365 y que "Cajamarca - Puente
     Cajamarca - Ibague" empieza ahi: la numeracion es continua de Armenia (PR 0) a
     Ibague (PR 80). [H] Los PR no miden exactamente 1 km (hay distancias > 1000 m).
  4. Tramos [H], alineados con las estaciones INVIAS del Act 1:
       Armenia - Calarca    km 0      a 5,0    (estacion 245, 5 km; FUERA del corredor)
       Paso Calarca-Cajam.  km 5,0    a 50,365 (estacion 244, 45 km)
       Cajamarca - Ibague   km 50,365 a 80,194 (estacion 243, 29 km)
     Los 5,0 km salen de la longitud INVIAS de la estacion 245, no del anexo.
  5. Cruza con los 8 siniestros fatales de prensa documentados en el Act 4 (+-1 dia).

SALIDAS (sin edad ni sexo de las victimas: solo agregados)
  microdato_ANSV_hechos.csv      una fila por hecho
  microdato_ANSV_por_km.csv      hechos y victimas por kilometro de abscisa
  cruce_prensa_ANSV.csv          los 8 hechos de prensa y si aparecen en el anexo
  microdato_ANSV_resumen.json    conteos que consumen la lamina y el informe
"""
import csv, json, os, sys, datetime as dt
from collections import Counter, defaultdict
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
ANEXO = os.path.join(BASE, "Sint_Via_LaLinea_CodTramo_4003_Vigencia_2021_Mar2026.xlsx")

LIM_CALARCA = 5.0        # [H] longitud INVIAS estacion 245
LIM_CAJAMARCA = 50.365   # [F] limite de sector declarado en el anexo
FIN_RUTA = 80.194        # [F] PR final del sector Cajamarca - Ibague en el anexo

def tramo(km):
    if km < LIM_CALARCA: return "Armenia – Calarcá (fuera)"
    if km < LIM_CAJAMARCA: return "Paso Calarcá – Cajamarca"
    return "Cajamarca – Ibagué"

# ---------- 1. lectura y autoverificacion contra las tablas de la ANSV ----------
wb = openpyxl.load_workbook(ANEXO, data_only=True)
ws = wb["Datos_ViaLaLinea_CodTramo_4003."] if "Datos_ViaLaLinea_CodTramo_4003." in wb.sheetnames else wb.worksheets[0]
filas = list(ws.iter_rows(values_only=True))
cab = filas[0]
V = [dict(zip(cab, f)) for f in filas[1:] if f[0] is not None]

def pivote(nombre):
    h = wb[nombre]; out = {}
    for f in h.iter_rows(min_row=2, values_only=True):
        if f[0] is None or f[0] == "Total Víctimas": continue
        out[f[0]] = (f[1] or 0, f[2] or 0)
    return out

errores = []
pv = pivote("Gravedad - Año")
for anio, (fal, les) in pv.items():
    f2 = sum(1 for v in V if v["Anio"] == anio and v["Gravedad"] == "Fallecido")
    l2 = sum(1 for v in V if v["Anio"] == anio and v["Gravedad"] == "Lesionado")
    if (f2, l2) != (fal, les): errores.append(f"Gravedad-Año {anio}: ANSV {fal}/{les}, registros {f2}/{l2}")
pu = pivote("Usuario via - Gravedad")
for u, (fal, les) in pu.items():
    f2 = sum(1 for v in V if v["UsuarioVia"] == u and v["Gravedad"] == "Fallecido")
    l2 = sum(1 for v in V if v["UsuarioVia"] == u and v["Gravedad"] == "Lesionado")
    if (f2, l2) != (fal, les): errores.append(f"Usuario {u}: ANSV {fal}/{les}, registros {f2}/{l2}")
if len({v["FID"] for v in V}) != len(V): errores.append("FID repetido")
if errores:
    print("ALTO: el anexo no reproduce sus propias tablas:", *errores, sep="\n  "); sys.exit(1)

# ---------- 2-4. hechos y abscisa ----------
for v in V:
    v["km"] = round(float(v["PR"]) + float(v["distancia"]) / 1000.0, 3)
    v["fecha"] = v["FechaHecho"].date()
    v["tramo"] = tramo(v["km"])

H = defaultdict(list)
for v in V: H[(v["fecha"], v["Municipio"], v["km"])].append(v)
hechos = []
for (fecha, mun, km), vs in sorted(H.items()):
    hechos.append({
        "fecha": fecha.isoformat(), "anio": fecha.year, "municipio": mun, "km": km,
        "tramo": vs[0]["tramo"], "sector_ansv": vs[0]["sector"],
        "lon": round(vs[0]["X"], 6), "lat": round(vs[0]["Y"], 6),
        "fallecidos": sum(1 for x in vs if x["Gravedad"] == "Fallecido"),
        "lesionados": sum(1 for x in vs if x["Gravedad"] == "Lesionado"),
        "usuario_predominante": Counter(x["UsuarioVia"] for x in vs).most_common(1)[0][0],
        "circunstancia": vs[0]["Circunstan"], "zona": vs[0]["Adm_Region"],
    })
with open(os.path.join(BASE, "microdato_ANSV_hechos.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(hechos[0])); w.writeheader(); w.writerows(hechos)

# por km (entero inferior de la abscisa)
K = defaultdict(lambda: {"hechos": 0, "hechos_fatales": 0, "fallecidos": 0, "lesionados": 0})
for h in hechos:
    k = int(h["km"]); K[k]["hechos"] += 1; K[k]["fallecidos"] += h["fallecidos"]
    K[k]["lesionados"] += h["lesionados"]; K[k]["hechos_fatales"] += 1 if h["fallecidos"] else 0
with open(os.path.join(BASE, "microdato_ANSV_por_km.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["km_desde", "km_hasta", "tramo", "hechos", "hechos_fatales", "fallecidos", "lesionados"])
    for k in range(0, int(FIN_RUTA) + 1):
        d = K[k]; w.writerow([k, k + 1, tramo(k + 0.5), d["hechos"], d["hechos_fatales"], d["fallecidos"], d["lesionados"]])

# ---------- 5. cruce con prensa (Act 4, ferropista/oe5-siniestralidad.md) ----------
PRENSA = [
    ("2022-03-31", "Villa del Cacique, tres tractomulas", "1", "Infobae", "no"),
    ("2022-09-22", "La Línea, chiva volcada", "2", "El Tiempo", "si"),
    ("2022-11-02", "La Línea, tractomula incendiada", "no indicado", "El Espectador", "si"),
    ("2023-12-14", "Descenso a Calarcá, bus volcado", "2", "El Tiempo", "si"),
    ("2024-05-22", "Alto de La Línea, tractomula", "1", "Infobae", "no"),
    ("2024-09-30", "Túnel Las Mariposas, auto con tres hermanos", "3", "El Tiempo", "si"),
    ("2025-05-16", "Túnel Los Azulejos, camión", "1", "El Tiempo", "si"),
    ("2025-05-24", "Puente helicoidal, bus universitario", "10-11", "El Tiempo / Infobae", "si"),
]
cruce = []
for fecha, lugar, muertos, medio, en_ponencia in PRENSA:
    f0 = dt.date.fromisoformat(fecha)
    m = [h for h in hechos if abs((dt.date.fromisoformat(h["fecha"]) - f0).days) <= 1 and h["fallecidos"] > 0]
    cruce.append({"fecha": fecha, "hecho": lugar, "muertos_prensa": muertos, "medio": medio,
                  "citado_en_ponencia": en_ponencia, "en_anexo_ANSV": "si" if m else "no",
                  "fallecidos_anexo": sum(h["fallecidos"] for h in m) if m else "",
                  "km_anexo": ";".join(str(h["km"]) for h in m)})
with open(os.path.join(BASE, "cruce_prensa_ANSV.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(cruce[0])); w.writeheader(); w.writerows(cruce)

# ---------- resumen ----------
def cuenta(filtro):
    hs = [h for h in hechos if filtro(h)]
    return {"hechos": len(hs), "hechos_fatales": sum(1 for h in hs if h["fallecidos"]),
            "fallecidos": sum(h["fallecidos"] for h in hs), "lesionados": sum(h["lesionados"] for h in hs)}
coords = Counter((h["lon"], h["lat"]) for h in hechos)
res = {
    "fuente": "ANSV-ONSV oficio 20265000140371 (22 sep 2026), anexo Ruta 4003, 2021 - mar 2026",
    "victimas": len(V), "fallecidos": sum(1 for v in V if v["Gravedad"] == "Fallecido"),
    "lesionados": sum(1 for v in V if v["Gravedad"] == "Lesionado"),
    "periodo": [min(v["fecha"] for v in V).isoformat(), max(v["fecha"] for v in V).isoformat()],
    "hechos": len(hechos), "puntos_distintos_victimas": len({(round(v["X"], 6), round(v["Y"], 6)) for v in V}),
    "puntos_distintos_hechos": len(coords),
    "punto_mas_cargado_hechos": {"lon_lat": list(coords.most_common(1)[0][0]), "hechos": coords.most_common(1)[0][1]},
    "por_tramo": {t: cuenta(lambda h, t=t: h["tramo"] == t) for t in sorted({h["tramo"] for h in hechos})},
    "paso_por_anio": {a: cuenta(lambda h, a=a: h["tramo"] == "Paso Calarcá – Cajamarca" and h["anio"] == a)
                      for a in range(2021, 2027)},
    "paso_km_11_a_41": cuenta(lambda h: 11 <= h["km"] < 41),
    "prensa": {"total": len(cruce), "en_anexo": sum(1 for c in cruce if c["en_anexo_ANSV"] == "si")},
    "limites_tramo_km": {"calarca": LIM_CALARCA, "cajamarca": LIM_CAJAMARCA, "fin": FIN_RUTA},
}
with open(os.path.join(BASE, "microdato_ANSV_resumen.json"), "w", encoding="utf-8") as f:
    json.dump(res, f, ensure_ascii=False, indent=2, sort_keys=True)
print(json.dumps(res, ensure_ascii=False, indent=1))
