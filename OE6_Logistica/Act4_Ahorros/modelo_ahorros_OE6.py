# -*- coding: utf-8 -*-
"""
OE 6 · Actividad 4 — Proyección de los ahorros operativos anuales para el sector transportador.
Producto: «Modelo de ahorros operativos» · Documento Excel. Sesión 19 (29 sep 2026).

Qué hace (lee los JSON de las Act 1, 2 y 3; ninguna cifra se copia a mano)
  1. Costo de un cruce por carretera, por sentido y escenario de tiempo (Act 1), con los unitarios del SICE-TAC (Act 3):
        C_carretera = (costo fijo/h × horas + combustible por terreno + variable/km × km + peajes) × (1 + f)
     Hipótesis: el combustible y el desgaste dependen de los km y del terreno, no de la demora [H].
  2. Costo del mismo cruce con Ferropista SIN la tarifa del servicio (que la fuente no publica):
        C_ferropista = (costo fijo/h × horas del ciclo Act 2 + km de acceso × (variable/km + combustible del terreno de acceso)) × (1 + f)
     Mientras el camión va sobre el tren no gasta combustible ni se desgasta, pero su costo fijo por hora sigue corriendo.
  3. Tarifa de indiferencia = C_carretera − C_ferropista: la tarifa máxima con la que al transportador le da igual usar la
     carretera o la Ferropista (principio de la dia. 19: «Peaje más bajo que los costos operacionales de la vía alternativa»).
     Es también el ahorro operativo BRUTO por viaje, a repartir entre el operador (tarifa) y el transportador.
  4. Proyección anual con la demanda de la fuente: 2.100 grandes camiones/día en el Alto de La Línea (INVÍAS 2017), 3 % anual,
     captación 90 % en 2030 (dia. 10); mitad por sentido [H]. Años 2026 y 2030, pesos constantes de sep 2026 [H].
  5. Sensibilidad a la tarifa (0, 50, 75 y 100 % de la tarifa de indiferencia) y contraste con las cifras de la fuente
     (5,0 M h/año ahorradas y $COP 2,5 billones de ahorro del sector transporte en 2030, dia. 28).
  6. Escribe modelo_ahorros_OE6.json (lo lee la Act 5) y Modelo_ahorros_operativos_OE6.xlsx.
Requiere openpyxl y pypdf.
"""
import os
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(AQUI))
import comun_OE6 as C  # noqa: E402
from openpyxl import Workbook  # noqa: E402

SALIDA_X = os.path.join(AQUI, "Modelo_ahorros_operativos_OE6.xlsx")
SALIDA_J = os.path.join(AQUI, "modelo_ahorros_OE6.json")

# [F, dia. 10] texto del PDF: «Grandes camiones : 2.100 veh./día», «Camiones medianos : 1.700 veh./día», «Tasa anual de crecimiento : 3%»,
# «Previsión de usuarios (año 2030): Grandes camiones : 90% · Camiones medianos : 60%»
TPD_GRANDES_2017, TPD_MEDIANOS_2017, CREC, CAPT_GRANDES, CAPT_MEDIANOS, ANO_BASE = 2100, 1700, 0.03, 0.90, 0.60, 2017
# [F, dia. 28] leído sobre la diapositiva renderizada: «Horas ahorradas 5,0 Mill. H./año» y «Ahorros $COP 2,5 Bill.» (sector transporte, año 2030)
FUENTE_HORAS_M, FUENTE_AHORRO_BILL, FUENTE_HORAS_DIA29_M = 5.0, 2.5, 5.3
ANOS = [2026, 2030]
REPARTO = 0.5          # [H] mitad del tránsito por sentido
TARIFAS = [0.0, 0.5, 0.75, 1.0]
PAR = {"Bajo": ("Mínimo", "Máximo"), "Central": ("Típico", "Típico"), "Alto": ("Máximo", "Mínimo")}


def main():
    C.texto_contiene("PON", ["Grandes camiones : 2.100 veh./día", "Camiones medianos : 1.700 veh./día", "Tasa anual de crecimiento : 3%",
                             "Grandes camiones : 90%", "Camiones medianos : 60%", "Peaje más bajo que los costos operacionales de la vía alternativa"])
    A1 = C.leer_json("Act1_TiemposRuta40", "registro_tiempos_ruta40_OE6.json")
    A2 = C.leer_json("Act2_TiempoFerropista", "tiempo_ciclo_ferropista_OE6.json")
    A3 = C.leer_json("Act3_CostosUnitarios", "matriz_costos_unitarios_OE6.json")
    u = A3["unitarios"]
    S = {"sentidos": {}, "rangos": list(PAR)}
    for k in C.RUTAS:
        r1, r2, ins = A1["sentidos"][k], A2["sentidos"][k], A3["sentidos"][k]["insumos"]
        comb = sum(r1["km"][t] * u["comb_km"][t] for t in C.TERRENOS)
        var_km = r1["km_total"] * u["cv_km_sin_comb"]
        peajes = ins["v_peajes"]
        out = {}
        for x, (er, ef) in PAR.items():
            c_road = (u["cf_hora"] * r1["horas"][er] + comb + var_km + peajes) * (1 + u["f_otros"])
            t_acc = A2["vel_acceso_terreno"][ef]
            c_fer = (u["cf_hora"] * r2["horas"][ef] + r2["km_acceso"] * (u["cv_km_sin_comb"] + u["comb_km"][t_acc])) * (1 + u["f_otros"])
            out[x] = {"h_carretera": r1["horas"][er], "h_ferropista": r2["horas"][ef], "dh": r1["horas"][er] - r2["horas"][ef],
                      "c_carretera": c_road, "c_ferropista_sin_tarifa": c_fer, "tarifa_indiferencia": c_road - c_fer,
                      "reduccion_pct_sin_tarifa": (c_road - c_fer) / c_road}
        S["sentidos"][k] = {"ruta": r1["ruta"], "combustible": comb, "variables_km": var_km, "peajes": peajes, "rangos": out}
    anual = {}
    for a in ANOS:
        g = TPD_GRANDES_2017 * (1 + CREC) ** (a - ANO_BASE)
        m = TPD_MEDIANOS_2017 * (1 + CREC) ** (a - ANO_BASE)
        cam = g * CAPT_GRANDES
        anual[a] = {"tpd_grandes": g, "tpd_medianos": m, "camiones_dia_captados": cam, "camiones_ano": cam * 365,
                    "ocupacion_capacidad": cam / A2["capacidad_camiones_dia"], "rangos": {}}
        for x in PAR:
            ah = sum(REPARTO * cam * 365 * S["sentidos"][k]["rangos"][x]["tarifa_indiferencia"] for k in C.RUTAS)
            hh = sum(REPARTO * cam * 365 * S["sentidos"][k]["rangos"][x]["dh"] for k in C.RUTAS)
            hm = sum(REPARTO * m * CAPT_MEDIANOS * 365 * S["sentidos"][k]["rangos"][x]["dh"] for k in C.RUTAS)
            anual[a]["rangos"][x] = {"ahorro_bruto": ah, "horas_grandes": hh, "horas_medianos_H": hm,
                                     "ahorro_transportador": {str(p): ah * (1 - p) for p in TARIFAS}}
    S["anual"] = anual
    S["contraste_fuente"] = {"horas_M_dia28": FUENTE_HORAS_M, "horas_M_dia29": FUENTE_HORAS_DIA29_M, "ahorro_bill_dia28": FUENTE_AHORRO_BILL}
    S["parametros"] = {"tpd_grandes_2017": TPD_GRANDES_2017, "tpd_medianos_2017": TPD_MEDIANOS_2017, "crecimiento": CREC,
                       "captacion_grandes": CAPT_GRANDES, "captacion_medianos": CAPT_MEDIANOS, "reparto_sentido": REPARTO, "tarifas": TARIFAS}
    C.guardar_json(SALIDA_J, S)
    excel(S, A1, A2, A3)
    for k, v in S["sentidos"].items():
        for x, y in v["rangos"].items():
            print(v["ruta"], x, f"dh {y['dh']:.3f} h  carretera {y['c_carretera']:,.0f}  ferro {y['c_ferropista_sin_tarifa']:,.0f}  T_ind {y['tarifa_indiferencia']:,.0f}  ({100 * y['reduccion_pct_sin_tarifa']:.1f} %)")
    for a, v in anual.items():
        print(a, f"camiones/día {v['camiones_dia_captados']:.0f}  ocupación {100 * v['ocupacion_capacidad']:.0f} %",
              {x: (f"{y['ahorro_bruto'] / 1e9:,.1f} mil M$", f"{y['horas_grandes'] / 1e6:.2f} M h", f"+med {y['horas_medianos_H'] / 1e6:.2f}") for x, y in v["rangos"].items()})


def excel(S, A1, A2, A3):
    wb = Workbook()
    C.portada(wb, "Modelo de ahorros operativos del sector transportador",
              "Objetivo específico 6 · Actividad 4 — Ahorro por cruce y proyección anual 2026 y 2030 (tractocamiones)",
              "4. Proyección de los ahorros operativos anuales para el sector transportador",
              "Modelo de ahorros operativos · Documento Excel",
              [("Entradas", "Tabla 1. Entradas leídas de las Act 1-3 y de la fuente primaria"),
               ("Por viaje", "Tabla 2. Costo de un cruce por carretera y por Ferropista (sin tarifa) y tarifa de indiferencia"),
               ("Anual", "Tabla 3. Demanda captada · Tabla 4. Ahorro bruto y horas ahorradas por año"),
               ("Tarifa", "Tabla 5. Ahorro que le queda al transportador según la tarifa"),
               ("Contraste", "Tabla 6. Contraste con las cifras de la fuente"),
               ("Fuentes", "Tabla 7. Fuentes citadas y huella md5")],
              "Índice de tablas: tablas 1 a 7 según el índice de hojas. Índice de figuras: no tiene (las figuras van en el informe de la Act 5). "
              "Índice de anexos: no tiene. Script: OE6_Logistica/Act4_Ahorros/modelo_ahorros_OE6.py.",
              "La fuente no publica la tarifa de la Ferropista: el ahorro bruto es un techo que se reparte entre operador y transportador. Solo "
              "tractocamiones con costos eficientes SICE-TAC; pesos constantes de sep 2026; demanda del proponente (INVÍAS 2017 proyectado).")
    u = A3["unitarios"]
    # ---- Entradas
    e = wb.create_sheet("Entradas"); e.sheet_view.showGridLines = False
    C.tit(e, "Tabla 1. Entradas del modelo")
    C.cab(e, 3, ["Entrada", "Valor", "Unidad", "Origen", "Marca"], [46, 16, 12, 52, 14])
    E = {}
    ent = [("cf_hora", "Costo fijo por hora de vehículo", u["cf_hora"], "$/h", "Act 3 · matriz_costos_unitarios_OE6.json", "[CP]", C.ENLACE),
           ("cv", "Costo variable por km sin combustible", u["cv_km_sin_comb"], "$/km", "Act 3", "[CP]", C.ENLACE),
           ("f", "Factor de otros costos f", u["f_otros"], "—", "Act 3", "[CP]", C.ENLACE)]
    ent += [("comb_" + t, f"Combustible por km · {t.lower()}", u["comb_km"][t], "$/km", "Act 3", "[CP]", C.ENLACE) for t in C.TERRENOS]
    ent += [("tpd_g", "Grandes camiones en el Alto de La Línea (2017)", S["parametros"]["tpd_grandes_2017"], "veh/día", "INVÍAS 2017 citado por la fuente", "[F, dia. 10]", C.ENTRADA),
            ("tpd_m", "Camiones medianos (2017)", S["parametros"]["tpd_medianos_2017"], "veh/día", "INVÍAS 2017 citado por la fuente", "[F, dia. 10]", C.ENTRADA),
            ("crec", "Crecimiento anual", S["parametros"]["crecimiento"], "—", "Fuente primaria", "[F, dia. 10]", C.ENTRADA),
            ("capg", "Captación de grandes camiones en 2030", S["parametros"]["captacion_grandes"], "—", "Fuente primaria", "[F, dia. 10]", C.ENTRADA),
            ("capm", "Captación de camiones medianos en 2030", S["parametros"]["captacion_medianos"], "—", "Fuente primaria", "[F, dia. 10]", C.ENTRADA),
            ("rep", "Fracción del tránsito por sentido", S["parametros"]["reparto_sentido"], "—", "Hipótesis", "[H]", C.NORMAL),
            ("capac", "Capacidad Ferropista (camiones/día)", A2["capacidad_camiones_dia"], "camiones/día", "Act 2 (140 trenes × 35)", "[F→CP]", C.ENLACE)]
    for i, (key, n, v, un, org, mc, fnt) in enumerate(ent):
        r = 4 + i
        E[key] = r
        fm = "0.0000%" if key == "f" else ("0.0%" if key in ("crec", "capg", "capm", "rep") else ("#,##0" if key in ("tpd_g", "tpd_m", "capac") else '"$"#,##0.00'))
        C.c(e, r, 1, n); C.c(e, r, 2, v, fnt, fm, h=(mc == "[H]")); C.c(e, r, 3, un); C.c(e, r, 4, org, wrap=True); C.c(e, r, 5, mc)
    r0 = 4 + len(ent) + 1
    C.cab(e, r0, ["Por sentido y escenario (Act 1 y Act 2)", A1["sentidos"]["IBG_CAL"]["ruta"], A1["sentidos"]["CAL_IBG"]["ruta"], "Unidad", "Marca"], col0=1)
    rr = r0 + 1
    for n, f, un in ([(f"Horas por carretera · {x}", lambda k, x=x: A1["sentidos"][k]["horas"][x], "h") for x in ["Mínimo", "Típico", "Máximo"]] +
                     [(f"Horas por Ferropista · {x}", lambda k, x=x: A2["sentidos"][k]["horas"][x], "h") for x in ["Mínimo", "Típico", "Máximo"]] +
                     [("Km por carretera", lambda k: A1["sentidos"][k]["km_total"], "km"),
                      ("Combustible del cruce por carretera (sin f)", lambda k: S["sentidos"][k]["combustible"], "$"),
                      ("Peajes del cruce por carretera", lambda k: S["sentidos"][k]["peajes"], "$"),
                      ("Km de acceso con Ferropista", lambda k: A2["sentidos"][k]["km_acceso"], "km")]):
        E[n] = rr
        C.c(e, rr, 1, n)
        for j, k in enumerate(C.RUTAS):
            C.c(e, rr, 2 + j, f(k), C.ENLACE, '"$"#,##0' if un == "$" else "0.000")
        C.c(e, rr, 4, un); C.c(e, rr, 5, "Act 1 / Act 2 / Act 3")
        rr += 1
    e.column_dimensions["C"].width = 18
    # ---- Por viaje
    p = wb.create_sheet("Por viaje"); p.sheet_view.showGridLines = False
    C.tit(p, "Tabla 2. Costo de un cruce Ibagué–Calarcá por carretera y por Ferropista (sin tarifa) y tarifa de indiferencia")
    C.cab(p, 3, ["Sentido", "Rango", "Horas carretera", "Horas Ferropista", "Horas ahorradas", "Costo carretera ($)", "Costo Ferropista sin tarifa ($)",
                 "Tarifa de indiferencia ($/camión)", "Reducción del costo (%)"], [20, 10, 12, 12, 12, 16, 18, 18, 14])
    col = {"IBG_CAL": "B", "CAL_IBG": "C"}
    acc_t = A2["vel_acceso_terreno"]
    r = 4
    PV = {}
    for k in C.RUTAS:
        X = col[k]
        for x, (er, ef) in PAR.items():
            C.c(p, r, 1, S["sentidos"][k]["ruta"]); C.c(p, r, 2, x, C.NEGRA)
            C.c(p, r, 3, f"=Entradas!{X}{E['Horas por carretera · ' + er]}", fmt="0.000")
            C.c(p, r, 4, f"=Entradas!{X}{E['Horas por Ferropista · ' + ef]}", fmt="0.000")
            C.c(p, r, 5, f"=C{r}-D{r}", fmt="0.000")
            C.c(p, r, 6, f"=(Entradas!$B${E['cf_hora']}*C{r}+Entradas!{X}{E['Combustible del cruce por carretera (sin f)']}+Entradas!{X}{E['Km por carretera']}*Entradas!$B${E['cv']}"
                         f"+Entradas!{X}{E['Peajes del cruce por carretera']})*(1+Entradas!$B${E['f']})", fmt='"$"#,##0')
            C.c(p, r, 7, f"=(Entradas!$B${E['cf_hora']}*D{r}+Entradas!{X}{E['Km de acceso con Ferropista']}*(Entradas!$B${E['cv']}+Entradas!$B${E['comb_' + acc_t[ef]]}))"
                         f"*(1+Entradas!$B${E['f']})", fmt='"$"#,##0')
            C.c(p, r, 8, f"=F{r}-G{r}", C.NEGRA, '"$"#,##0', fill=C.FILL_RES); C.c(p, r, 9, f"=H{r}/F{r}", fmt=C.FMT_PCT)
            PV[(k, x)] = r
            r += 1
        r += 1
    C.nota(p, r, "Rango bajo = carretera en flujo libre (modelo SICE-TAC) frente a la Ferropista más lenta; central = típico frente a típico; alto = carretera con "
                 "pare y siga frente a la Ferropista más rápida. Una tarifa de indiferencia negativa significa que, en ese caso, la carretera es más barata "
                 "aunque la Ferropista no cobre nada. No incluye cargue ni descargue de la mercancía (iguales en los dos modos).", ancho=9)
    # ---- Anual
    an = wb.create_sheet("Anual"); an.sheet_view.showGridLines = False
    C.tit(an, "Tabla 3. Demanda captada por la Ferropista (tractocamiones)")
    C.cab(an, 3, ["Año", "Grandes camiones/día (proyección)", "Captación", "Camiones/día captados", "Camiones/año", "Ocupación de la capacidad"], [10, 18, 12, 16, 16, 16])
    FA = {}
    for i, a in enumerate(ANOS):
        r = 4 + i
        FA[a] = r
        C.c(an, r, 1, a, C.NEGRA); C.c(an, r, 2, f"=Entradas!$B${E['tpd_g']}*(1+Entradas!$B${E['crec']})^(A{r}-{ANO_BASE})", fmt="#,##0")
        C.c(an, r, 3, f"=Entradas!$B${E['capg']}", fmt="0%"); C.c(an, r, 4, f"=B{r}*C{r}", fmt="#,##0"); C.c(an, r, 5, f"=D{r}*365", fmt="#,##0")
        C.c(an, r, 6, f"=D{r}/Entradas!$B${E['capac']}", fmt="0%")
    C.nota(an, 7, "La captación del 90 % es la prevista por la fuente para 2030; se aplica igual a 2026 para mostrar el orden de magnitud con la demanda de hoy [H].", ancho=6)
    an.cell(9, 1, "Tabla 4. Ahorro operativo bruto anual y horas ahorradas").font = C.Font(name=C.F, size=13, bold=True, color=C.AZUL)
    C.cab(an, 10, ["Año", "Rango", "Ahorro bruto ($/año)", "Ahorro bruto (miles de millones $)", "Horas ahorradas, grandes camiones", "Horas ahorradas, medianos [H]", "Horas totales (millones)"])
    an.column_dimensions["G"].width = 16
    an.column_dimensions["C"].width = 20
    r = 11
    AR = {}
    for a in ANOS:
        for x in PAR:
            fa = FA[a]
            C.c(an, r, 1, a, C.NEGRA); C.c(an, r, 2, x, C.NEGRA)
            C.c(an, r, 3, f"=Entradas!$B${E['rep']}*$E${fa}*('Por viaje'!H{PV[('IBG_CAL', x)]}+'Por viaje'!H{PV[('CAL_IBG', x)]})", fmt='"$"#,##0', fill=C.FILL_RES)
            C.c(an, r, 4, f"=C{r}/1E9", fmt="#,##0.0")
            C.c(an, r, 5, f"=Entradas!$B${E['rep']}*$E${fa}*('Por viaje'!E{PV[('IBG_CAL', x)]}+'Por viaje'!E{PV[('CAL_IBG', x)]})", fmt="#,##0")
            C.c(an, r, 6, f"=Entradas!$B${E['rep']}*Entradas!$B${E['tpd_m']}*(1+Entradas!$B${E['crec']})^(A{r}-{ANO_BASE})*Entradas!$B${E['capm']}*365"
                          f"*('Por viaje'!E{PV[('IBG_CAL', x)]}+'Por viaje'!E{PV[('CAL_IBG', x)]})", fmt="#,##0", h=True)
            C.c(an, r, 7, f"=(E{r}+F{r})/1E6", fmt="0.00")
            AR[(a, x)] = r
            r += 1
    C.nota(an, r + 1, "Horas de camiones medianos [H]: se les aplica la misma reducción de tiempo que a los tractocamiones solo para comparar con la cifra de "
                      "horas de la fuente; su ahorro en pesos no se calcula porque el SICE-TAC se consultó solo para 3S3.", ancho=7)
    # ---- Tarifa
    t = wb.create_sheet("Tarifa"); t.sheet_view.showGridLines = False
    C.tit(t, "Tabla 5. Ahorro anual que le queda al transportador según la tarifa (miles de millones de $, año 2030)")
    C.cab(t, 3, ["Tarifa como fracción de la tarifa de indiferencia"] + list(PAR), [30, 16, 16, 16])
    for i, pct in enumerate(TARIFAS):
        r = 4 + i
        C.c(t, r, 1, pct, C.ENTRADA, "0%")
        for j, x in enumerate(PAR):
            C.c(t, r, 2 + j, f"=Anual!D{AR[(2030, x)]}*(1-$A{r})", fmt="#,##0.0")
    C.nota(t, 9, "Con tarifa = 100 % de la de indiferencia, todo el ahorro lo captura el operador y el transportador queda igual que por carretera. "
                 "La fuente no dice cómo se reparte.", ancho=4)
    # ---- Contraste
    cs = wb.create_sheet("Contraste"); cs.sheet_view.showGridLines = False
    C.tit(cs, "Tabla 6. Contraste con las cifras de la fuente (año 2030)")
    C.cab(cs, 3, ["Indicador", "Fuente", "Este modelo · bajo", "Este modelo · central", "Este modelo · alto", "Marca de la fuente"], [44, 14, 16, 16, 16, 22])
    C.c(cs, 4, 1, "Horas ahorradas al año (millones)"); C.c(cs, 4, 2, FUENTE_HORAS_M, C.ENTRADA, "0.0")
    for j, x in enumerate(PAR):
        C.c(cs, 4, 3 + j, f"=Anual!G{AR[(2030, x)]}", fmt="0.00")
    C.c(cs, 4, 6, "[F, dia. 28]; la dia. 29 rotula 5,3")
    C.c(cs, 5, 1, "Ahorro del sector transporte (billones de $)"); C.c(cs, 5, 2, FUENTE_AHORRO_BILL, C.ENTRADA, "0.0")
    for j, x in enumerate(PAR):
        C.c(cs, 5, 3 + j, f"=Anual!D{AR[(2030, x)]}/1000", fmt="0.000")
    C.c(cs, 5, 6, "[F, dia. 28]")
    C.c(cs, 6, 1, "Tiempo de cruce: actual → Ferropista"); C.c(cs, 6, 2, "4 h → 70 min", C.ENTRADA)
    C.c(cs, 6, 3, "ver Act 2")
    C.c(cs, 6, 4, f"{C.es(A1['sentidos']['IBG_CAL']['horas']['Típico'], 2)} h → {C.es(A2['sentidos']['IBG_CAL']['horas']['Típico'] * 60, 0)} min (Ibagué → Calarcá)", C.ENLACE)
    C.c(cs, 6, 6, "[F, dia. 29]")
    C.nota(cs, 8, "La fuente no explica cómo obtiene 5,0 M h ni 2,5 billones (qué vehículos, qué valor del tiempo, qué año de precios). Este modelo cuenta "
                  "solo costos de operación de tractocamiones con precios de sep 2026; no incluye el valor del tiempo de la carga, ni de los pasajeros, ni "
                  "externalidades. Por eso la diferencia no prueba que la fuente esté mal: muestra cuánto de su cifra no se puede reproducir con datos "
                  "oficiales de costo de operación.", ancho=6)
    C.hoja_fuentes(wb, ["PON", "SICETAC_IBG_CAL", "SICETAC_CAL_IBG"], 7)
    for w in wb.worksheets:
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(SALIDA_X)


if __name__ == "__main__":
    main()
