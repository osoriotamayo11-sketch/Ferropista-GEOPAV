# -*- coding: utf-8 -*-
"""
OE 6 · Actividad 2 — Cálculo del tiempo de ciclo del sistema Ferropista (carga, desplazamiento y descarga).
Producto: «Memoria de tiempos logísticos» · Documento Excel. Sesión 19 (29 sep 2026).

Qué hace
  1. Toma de la fuente primaria el ciclo de operación: peaje y carga 25 min + desplazamiento 30 min + descarga 15 min
     = 70 min (dia. 23, leído sobre la diapositiva renderizada) y la frecuencia de la fase inicial, 140 trenes/día (dia. 20).
  2. Añade lo que la ponencia no cuenta y el camión sí vive, para comparar cabecera con cabecera como en la Act 1:
       - acceso por carretera cabecera → terminal y terminal → cabecera. Hipótesis: las terminales están en los portales
         adoptados por el OE 1 [H]. Distancia vial de OpenStreetMap (OSRM) [F→CP] y velocidad del SICE-TAC por escenario:
         mínimo = ondulado 33,13 km/h, típico = montaña/urbano 23,57 km/h, máximo = afirmado 15 km/h [F SICE-TAC, asignación H].
       - espera del siguiente tren: 140 trenes/día repartidos por igual entre los dos sentidos y 24 h de operación [H]
         → intervalo de 20,6 min; espera mínima 0, típica medio intervalo, máxima un intervalo [CP].
  3. Compara con los tiempos de la Act 1 (lee su JSON) y arma el rango de reducción:
       bajo = carretera mínima − Ferropista máxima · central = típica − típica · alto = carretera máxima − Ferropista mínima.
  4. Escribe tiempo_ciclo_ferropista_OE6.json (lo lee la Act 4) y Memoria_tiempos_logisticos_OE6.xlsx.
Requiere openpyxl y pypdf.
"""
import json
import os
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(AQUI))
import comun_OE6 as C  # noqa: E402
from openpyxl import Workbook  # noqa: E402

SALIDA_X = os.path.join(AQUI, "Memoria_tiempos_logisticos_OE6.xlsx")
SALIDA_J = os.path.join(AQUI, "tiempo_ciclo_ferropista_OE6.json")
DATOS = os.path.join(AQUI, "Datos")

# [F, dia. 23] leído sobre la diapositiva renderizada (la cifra está en una imagen, no en el texto del PDF)
CICLO_MIN = {"Peaje y carga": 25.0, "Desplazamiento": 30.0, "Descarga": 15.0}
# [F, dia. 20] «Fase inicial: 140 trenes/día» · «Hasta 35 tractomulas por tren»
TRENES_DIA, CAMIONES_TREN = 140, 35
HORAS_OPERACION = 24.0          # [H]
REPARTO_SENTIDO = 0.5           # [H] mitad de los trenes en cada sentido
VEL_ACCESO = {"Mínimo": "Ondulado", "Típico": "Montaña", "Máximo": "Afirmado"}   # [H] asignación de la velocidad SICE-TAC
LEGS = {"IBG_CAL": ("ibague_portalE", "portalW_calarca"), "CAL_IBG": ("calarca_portalW", "portalE_ibague")}


def osrm(n):
    d = json.load(open(os.path.join(DATOS, f"osrm_{n}.json"), encoding="utf-8"))
    r = d["routes"][0]
    return {"km": r["distance"] / 1000, "min_auto": r["duration"] / 60,
            "snap_m": [w["distance"] for w in d["waypoints"]], "via": [w["name"] for w in d["waypoints"]]}


def main():
    C.texto_contiene("PON", ["Fase inicial: 140 trenes/día", "Hasta 35 tractomulas por tren", "Ferropista: 70 min (tiempo integrado)"])
    A1 = C.leer_json("Act1_TiemposRuta40", "registro_tiempos_ruta40_OE6.json")
    vel = C.leer_sicetac("IBG_CAL")["vel_kmh"]
    oe1 = json.load(open(os.path.join(C.RAIZ, "OE1_Topografia", "cifras_OE1.json"), encoding="utf-8"))
    intervalo_min = HORAS_OPERACION * 60 / (TRENES_DIA * REPARTO_SENTIDO)
    espera = {"Mínimo": 0.0, "Típico": intervalo_min / 2, "Máximo": intervalo_min}
    op_min = sum(CICLO_MIN.values())
    assert op_min == 70.0
    S = {"ciclo_operacion_min": CICLO_MIN, "ciclo_operacion_total_min": op_min, "trenes_dia": TRENES_DIA, "camiones_tren": CAMIONES_TREN,
         "horas_operacion": HORAS_OPERACION, "reparto_sentido": REPARTO_SENTIDO, "intervalo_min": intervalo_min, "espera_min": espera,
         "vel_acceso_kmh": {e: vel[t] for e, t in VEL_ACCESO.items()}, "vel_acceso_terreno": VEL_ACCESO,
         "portales": {"oriental": oe1["portal_oriental"], "occidental": oe1["portal_occidental"]},
         "capacidad_camiones_dia": TRENES_DIA * CAMIONES_TREN, "sentidos": {}}
    for k, (a, b) in LEGS.items():
        la, lb = osrm(a), osrm(b)
        km_acc = la["km"] + lb["km"]
        horas = {}
        for e in ["Mínimo", "Típico", "Máximo"]:
            t_acc = km_acc / vel[VEL_ACCESO[e]] * 60
            horas[e] = (t_acc + espera[e] + op_min) / 60
        road = A1["sentidos"][k]["horas"]
        red = {"Bajo": road["Mínimo"] - horas["Máximo"], "Central": road["Típico"] - horas["Típico"], "Alto": road["Máximo"] - horas["Mínimo"]}
        base = {"Bajo": road["Mínimo"], "Central": road["Típico"], "Alto": road["Máximo"]}
        S["sentidos"][k] = {"ruta": A1["sentidos"][k]["ruta"], "acceso_origen": {"tramo": a, **la}, "acceso_destino": {"tramo": b, **lb},
                            "km_acceso": km_acc, "horas": horas, "horas_carretera": road, "reduccion_h": red,
                            "reduccion_pct": {x: red[x] / base[x] for x in red}}
    S["osrm_via"] = {k: osrm(n) for k, n in {"IBG_CAL": "ibague_calarca_via", "CAL_IBG": "calarca_ibague_via"}.items()}
    S["md5"] = {f: C.md5(os.path.join(DATOS, f)) for f in sorted(os.listdir(DATOS))}
    C.guardar_json(SALIDA_J, S)
    excel(S, A1)
    for k, v in S["sentidos"].items():
        print(v["ruta"], "acceso km", round(v["km_acceso"], 2), {e: round(h, 3) for e, h in v["horas"].items()},
              "reducción h", {x: round(y, 2) for x, y in v["reduccion_h"].items()}, {x: f"{100 * y:.0f}%" for x, y in v["reduccion_pct"].items()})


def excel(S, A1):
    wb = Workbook()
    C.portada(wb, "Memoria de tiempos logísticos del sistema Ferropista",
              "Objetivo específico 6 · Actividad 2 — Tiempo de ciclo cabecera–cabecera con carga, desplazamiento y descarga",
              "2. Cálculo del tiempo de ciclo del sistema Ferropista (carga, desplazamiento y descarga)",
              "Memoria de tiempos logísticos · Documento Excel",
              [("Parámetros", "Tabla 1. Ciclo de operación de la fuente, frecuencia y velocidades de acceso"),
               ("Accesos", "Tabla 2. Accesos por carretera entre cabeceras y terminales (OpenStreetMap)"),
               ("Ciclo", "Tabla 3. Tiempo de ciclo Ferropista por sentido y escenario"),
               ("Comparación", "Tabla 4. Reducción del tiempo de ciclo frente a la Ruta 40 (Act 1)"),
               ("Fuentes", "Tabla 5. Fuentes citadas y huella md5")],
              "Índice de tablas: tablas 1 a 5 según el índice de hojas. Índice de figuras: no tiene. Índice de anexos: respuestas crudas de OSRM y "
              "Nominatim en Act2_TiempoFerropista/Datos/. Script: OE6_Logistica/Act2_TiempoFerropista/tiempo_ciclo_ferropista_OE6.py.",
              "La ubicación de las terminales no la da la fuente: se suponen en los portales del OE 1 [H]. La fuente tampoco dice cómo se reparten los "
              "140 trenes ni las horas de servicio. Los 70 min son una cifra del proponente, no auditada.")
    # ---- Parámetros
    p = wb.create_sheet("Parámetros"); p.sheet_view.showGridLines = False
    C.tit(p, "Tabla 1. Parámetros del ciclo Ferropista")
    C.cab(p, 3, ["Parámetro", "Valor", "Unidad", "Marca / fuente"], [44, 12, 12, 60])
    filas = [("Peaje y carga", S["ciclo_operacion_min"]["Peaje y carga"], "min", "[F, dia. 23]", True),
             ("Desplazamiento", S["ciclo_operacion_min"]["Desplazamiento"], "min", "[F, dia. 23]", True),
             ("Descarga", S["ciclo_operacion_min"]["Descarga"], "min", "[F, dia. 23]", True),
             ("Ciclo de operación (suma)", "=SUM(B4:B6)", "min", "[CP] (la fuente publica 70 min: control)", False),
             ("Trenes por día, fase inicial", S["trenes_dia"], "trenes/día", "[F, dia. 20]", True),
             ("Camiones por tren", S["camiones_tren"], "camiones", "[F, dia. 20] «hasta 35»", True),
             ("Horas de operación al día", S["horas_operacion"], "h", "[H] servicio continuo", "H"),
             ("Fracción de trenes por sentido", S["reparto_sentido"], "—", "[H] reparto por igual", "H"),
             ("Intervalo entre trenes de un sentido", "=B10*60/(B8*B11)", "min", "[CP]", False),
             ("Capacidad (camiones/día, dos sentidos)", "=B8*B9", "camiones/día", "[CP]", False),
             ("Velocidad de acceso · mínimo (ondulado)", S["vel_acceso_kmh"]["Mínimo"], "km/h", "[F] SICE-TAC · asignación [H]", True),
             ("Velocidad de acceso · típico (montaña/urbano)", S["vel_acceso_kmh"]["Típico"], "km/h", "[F] SICE-TAC · asignación [H]", True),
             ("Velocidad de acceso · máximo (afirmado)", S["vel_acceso_kmh"]["Máximo"], "km/h", "[F] SICE-TAC · asignación [H]", True)]
    for i, (a, v, u, m, ent) in enumerate(filas):
        r = 4 + i
        C.c(p, r, 1, a); C.c(p, r, 2, v, C.ENTRADA if ent is True else C.NORMAL, "0.00", h=(ent == "H")); C.c(p, r, 3, u); C.c(p, r, 4, m)
    C.nota(p, 18, "Fila 7 debe dar 70 min, el «tiempo integrado» que la fuente publica en la dia. 29. Espera en terminal: mínimo 0, típico medio "
                  "intervalo (llegada al azar), máximo un intervalo. Las velocidades de acceso son las del SICE-TAC para 3S3 cargado; qué terreno "
                  "representa cada escenario es hipótesis.", ancho=4)
    # ---- Accesos
    a = wb.create_sheet("Accesos"); a.sheet_view.showGridLines = False
    C.tit(a, "Tabla 2. Accesos por carretera entre las cabeceras municipales y las terminales supuestas en los portales del OE 1")
    C.cab(a, 3, ["Sentido", "Tramo", "Distancia vial (km)", "Tiempo en auto OSRM (min)", "Ajuste al vial en origen / destino (m)", "Vía del ajuste"], [20, 22, 14, 14, 20, 34])
    r = 4
    fila_acc = {}
    for k, v in S["sentidos"].items():
        fila_acc[k] = r
        for leg in ("acceso_origen", "acceso_destino"):
            L = v[leg]
            C.c(a, r, 1, v["ruta"]); C.c(a, r, 2, L["tramo"]); C.c(a, r, 3, L["km"], C.ENTRADA, "0.000")
            C.c(a, r, 4, round(L["min_auto"], 1), C.ENTRADA, "0.0"); C.c(a, r, 5, " / ".join(C.es(x, 0) for x in L["snap_m"]))
            C.c(a, r, 6, " / ".join(x or "(sin nombre)" for x in L["via"]))
            r += 1
        C.c(a, r, 1, v["ruta"], C.NEGRA); C.c(a, r, 2, "Total acceso", C.NEGRA); C.c(a, r, 3, f"=C{r - 2}+C{r - 1}", C.NEGRA, "0.00", fill=C.FILL_RES)
        fila_acc[k] = r
        r += 2
    C.nota(a, r, f"Portal oriental (OE 1): {S['portales']['oriental']['lat']}, {S['portales']['oriental']['lon']}; se ajusta a la «Vía a la Montañita», una vía rural. "
                 f"Portal occidental: {S['portales']['occidental']['lat']}, {S['portales']['occidental']['lon']}; queda a unos 258 m del vial más cercano. "
                 "Cabeceras: centro administrativo de OpenStreetMap (Nominatim). Una terminal de 35 tractomulas por tren exigiría vías de acceso "
                 "nuevas o mejoradas: no se evalúa aquí.", ancho=6)
    r += 2
    C.nota(a, r, "Control: OSRM da entre cabeceras por la Ruta 40 " + " y ".join(f"{C.es(v['km'], 1)} km ({C.es(v['min_auto'], 0)} min en auto)" for v in S["osrm_via"].values())
           + " para Ibagué→Calarcá y Calarcá→Ibagué; la misma asimetría que el SICE-TAC (85,01 frente a 76,17 km).", ancho=6)
    # ---- Ciclo
    c = wb.create_sheet("Ciclo"); c.sheet_view.showGridLines = False
    C.tit(c, "Tabla 3. Tiempo de ciclo Ferropista cabecera–cabecera por sentido y escenario")
    C.cab(c, 3, ["Sentido", "Escenario", "Acceso (min)", "Espera del tren (min)", "Operación (min)", "Total (min)", "Total (h)"], [20, 12, 13, 13, 13, 13, 11])
    r = 4
    vfila = {"Mínimo": 14, "Típico": 15, "Máximo": 16}
    esp = {"Mínimo": "0", "Típico": "Parámetros!B12/2", "Máximo": "Parámetros!B12"}
    fila_ciclo = {}
    for k, v in S["sentidos"].items():
        for e in ["Mínimo", "Típico", "Máximo"]:
            C.c(c, r, 1, v["ruta"]); C.c(c, r, 2, e, C.NEGRA)
            C.c(c, r, 3, f"=Accesos!C{fila_acc[k]}/Parámetros!B{vfila[e]}*60", fmt="0.0")
            C.c(c, r, 4, f"={esp[e]}", fmt="0.0"); C.c(c, r, 5, "=Parámetros!B7", fmt="0.0")
            C.c(c, r, 6, f"=C{r}+D{r}+E{r}", C.NEGRA, "0.0", fill=C.FILL_RES); C.c(c, r, 7, f"=F{r}/60", fmt="0.000")
            fila_ciclo[(k, e)] = r
            r += 1
        r += 1
    # ---- Comparación
    m = wb.create_sheet("Comparación"); m.sheet_view.showGridLines = False
    C.tit(m, "Tabla 4. Reducción del tiempo de ciclo frente a la operación actual por la Ruta 40")
    C.cab(m, 3, ["Sentido", "Rango", "Carretera (h) · Act 1", "Ferropista (h)", "Reducción (h)", "Reducción (%)", "Emparejamiento"], [20, 10, 16, 14, 13, 13, 44])
    r = 4
    par = {"Bajo": ("Mínimo", "Máximo"), "Central": ("Típico", "Típico"), "Alto": ("Máximo", "Mínimo")}
    for k, v in S["sentidos"].items():
        for x, (er, ef) in par.items():
            C.c(m, r, 1, v["ruta"]); C.c(m, r, 2, x, C.NEGRA)
            C.c(m, r, 3, A1["sentidos"][k]["horas"][er], C.ENLACE, "0.000"); C.c(m, r, 4, f"=Ciclo!G{fila_ciclo[(k, ef)]}", fmt="0.000")
            C.c(m, r, 5, f"=C{r}-D{r}", C.NEGRA, "0.000", fill=C.FILL_RES); C.c(m, r, 6, f"=E{r}/C{r}", fmt=C.FMT_PCT)
            C.c(m, r, 7, f"Carretera {er.lower()} − Ferropista {ef.lower()}")
            r += 1
        r += 1
    C.nota(m, r, "Columna C en verde: valor leído del JSON de la Act 1 (registro_tiempos_ruta40_OE6.json). La fuente compara 4 h con 70 min (−71 %) "
                 "sin contar accesos ni espera del tren; esta tabla sí los cuenta.", ancho=7)
    C.hoja_fuentes(wb, ["PON", "SICETAC_IBG_CAL", "OSRM", "OE1"], 5)
    for w in wb.worksheets:
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(SALIDA_X)


if __name__ == "__main__":
    main()
