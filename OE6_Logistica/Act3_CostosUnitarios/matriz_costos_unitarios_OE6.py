# -*- coding: utf-8 -*-
"""
OE 6 · Actividad 3 — Estructuración de la matriz de costos operativos unitarios del transporte de carga.
Producto: «Matriz de costos unitarios» · Documento Excel (verificador). Sesión 19 (29 sep 2026).

Qué hace
  1. Lee las dos consultas oficiales SICE-TAC guardadas en PDF (3S3 portacontenedores, contenedor cargado, periodo
     20260901): costos fijos del mes, horas hábiles, parámetros por kilómetro (llantas, lubricantes, filtros, mantenimiento,
     lavado), rendimiento de combustible por terreno, precio del ACPM, tasas de imprevistos y de «otros costos».
     Peajes: listado oficial «Peajes por rutas con tarifas» (corte 01-09-2026), categoría V.
  2. Descompone el costo en unidades que sirven para comparar carretera y Ferropista [CP]:
        costo fijo por hora de vehículo = costos fijos mes / horas hábiles mes
        costo variable por km (sin combustible) = Σ parámetros por km × (1 + imprevistos)
        combustible por km y terreno = ACPM / rendimiento
        factor de otros costos f = (comisiones + administrativo + retefuente e ICA) / (fijos + variables)
  3. Reproduce con esas unidades el «costo de movilización» y el «costo total del viaje» que publica el SICE-TAC en los dos
     sentidos. Si la diferencia pasa del 0,1 %, el script se detiene.
  4. Escribe matriz_costos_unitarios_OE6.json (lo lee la Act 4) y Matriz_costos_unitarios_OE6.xlsx.
Pesos corrientes de septiembre de 2026. Requiere openpyxl y pypdf.
"""
import os
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(AQUI))
import comun_OE6 as C  # noqa: E402
from openpyxl import Workbook  # noqa: E402

SALIDA_X = os.path.join(AQUI, "Matriz_costos_unitarios_OE6.xlsx")
SALIDA_J = os.path.join(AQUI, "matriz_costos_unitarios_OE6.json")
KM_ITEMS = ["km_llantas", "km_lubricantes", "km_filtros", "km_mantenimiento", "km_lavado"]
TOL = 0.001


def unitarios(d):
    base = d["cf_viaje"] + d["cv_viaje"]
    u = {"cf_hora": d["cf_mes"] / d["horas_habiles_mes"],
         "cv_km_sin_comb": sum(d[x] for x in KM_ITEMS) * (1 + d["tasa_imprevistos"]),
         "comb_km": {t: d["acpm"] / d["rend_kmgal"][t] for t in C.TERRENOS},
         "f_otros": (d["o_comisiones"] + d["o_admin"] + d["o_rete"]) / base,
         "f_detalle": {"comisiones_prestacional": d["o_comisiones"] / base, "administrativo": d["o_admin"] / base, "retefuente_ica": d["o_rete"] / base}}
    u["costo_hora_total"] = u["cf_hora"] * (1 + u["f_otros"])
    return u


def reproducir(d, u):
    h = sum(d["dist_km"][t] / d["vel_kmh"][t] for t in C.TERRENOS)
    comb = sum(d["dist_km"][t] * u["comb_km"][t] for t in C.TERRENOS)
    cv = comb + d["v_peajes"] + d["dist_total_km"] * u["cv_km_sin_comb"]
    mov = (u["cf_hora"] * h + cv) * (1 + u["f_otros"])
    total = (u["cf_hora"] * (h + d["horas_espera"]) + cv) * (1 + u["f_otros"])
    return {"horas": h, "combustible": comb, "variables": cv, "movilizacion": mov, "total": total}


def main():
    C.texto_contiene("SICETAC_IBG_CAL", ["Las horas operativas al mes pasan de 238 horas", "a 230 horas (182 horas ordinarias más 48 horas extra)"])
    S = {"sentidos": {}}
    for k in C.RUTAS:
        d = C.leer_sicetac(k)
        u = unitarios(d)
        rp = reproducir(d, u)
        e_mov = rp["movilizacion"] / d["costo_mov_carga"] - 1
        e_tot = rp["total"] / d["costo_total_viaje"] - 1
        e_hora = u["costo_hora_total"] / d["costo_hora"] - 1
        for n, e in [("movilización", e_mov), ("total", e_tot), ("costo hora", e_hora)]:
            if abs(e) > TOL:
                sys.exit(f"ALTO: la reproducción del {n} SICE-TAC ({k}) difiere {100 * e:.3f} %")
        S["sentidos"][k] = {"ruta": d["ruta"], "insumos": d, "unitarios": u, "reproduccion": rp,
                            "error_rel": {"movilizacion": e_mov, "total": e_tot, "costo_hora": e_hora}}
    # los parámetros unitarios deben ser idénticos en los dos sentidos (mismo vehículo, mismo periodo)
    a, b = (S["sentidos"][k]["unitarios"] for k in C.RUTAS)
    for x in ["cf_hora", "cv_km_sin_comb", "costo_hora_total"]:
        assert abs(a[x] / b[x] - 1) < 1e-6, x
    assert abs(a["f_otros"] - b["f_otros"]) < 1e-5
    S["unitarios"] = a
    S["peajes"] = {k: S["sentidos"][k]["insumos"]["peajes_lista"] for k in C.RUTAS}
    S["md5"] = {k: C.md5(C.ARCHIVOS[k]) for k in ["SICETAC_IBG_CAL", "SICETAC_CAL_IBG", "SICETAC_PEAJ", "RES2024"]}
    C.guardar_json(SALIDA_J, S)
    excel(S)
    u = S["unitarios"]
    print(f"costo fijo/h {u['cf_hora']:.2f}  costo hora con otros {u['costo_hora_total']:.2f}  cv/km {u['cv_km_sin_comb']:.2f}  f {u['f_otros']:.6f}")
    print({t: round(v, 2) for t, v in u["comb_km"].items()})
    for k, v in S["sentidos"].items():
        print(v["ruta"], {x: round(y, 2) for x, y in v["reproduccion"].items()}, {x: f"{100 * y:+.4f}%" for x, y in v["error_rel"].items()})


def excel(S):
    wb = Workbook()
    C.portada(wb, "Matriz de costos operativos unitarios del transporte de carga",
              "Objetivo específico 6 · Actividad 3 — Tractocamión 3S3 portacontenedor, costos eficientes SICE-TAC (sep 2026)",
              "3. Estructuración de la matriz de costos operativos unitarios del transporte de carga",
              "Matriz de costos unitarios · Documento Excel",
              [("Insumos", "Tabla 1. Insumos de las consultas SICE-TAC por sentido"),
               ("Unitarios", "Tabla 2. Costos unitarios por hora y por kilómetro · Tabla 3. Combustible por terreno · Tabla 4. Factor de otros costos"),
               ("Verificación", "Tabla 5. Reproducción del costo SICE-TAC con los unitarios"),
               ("Peajes", "Tabla 6. Peajes de la Ruta 40 por sentido, categoría V"),
               ("Fuentes", "Tabla 7. Fuentes citadas y huella md5")],
              "Índice de tablas: tablas 1 a 7 según el índice de hojas. Índice de figuras: no tiene. Índice de anexos: las consultas SICE-TAC en PDF, "
              "en OE6_Logistica/Fuentes/. Script: OE6_Logistica/Act3_CostosUnitarios/matriz_costos_unitarios_OE6.py.",
              "Son costos EFICIENTES de referencia del Ministerio (piso para la relación transportador–generador), no costos observados de una flota. "
              "Un solo vehículo representativo (3S3 portacontenedor); los camiones medianos no se modelan.")
    k1, k2 = list(C.RUTAS)
    s1, s2 = S["sentidos"][k1]["insumos"], S["sentidos"][k2]["insumos"]
    # ---- Insumos
    ws = wb.create_sheet("Insumos"); ws.sheet_view.showGridLines = False
    C.tit(ws, "Tabla 1. Insumos leídos de las consultas SICE-TAC (periodo 20260901)")
    C.cab(ws, 3, ["Insumo", s1["ruta"], s2["ruta"], "Unidad"], [46, 18, 18, 14])
    filas = [("Costos fijos del mes", "cf_mes", "$/mes"), ("Horas hábiles del mes", "horas_habiles_mes", "h"), ("Horas logísticas (cargue + descargue)", "horas_espera", "h"),
             ("Valor ACPM", "acpm", "$/galón"), ("Llantas", "km_llantas", "$/km"), ("Lubricantes", "km_lubricantes", "$/km"), ("Filtros", "km_filtros", "$/km"),
             ("Mantenimiento y reparación", "km_mantenimiento", "$/km"), ("Lavado y engrase", "km_lavado", "$/km"), ("Tasa de imprevistos", "tasa_imprevistos", "—"),
             ("Subtotal costos fijos del viaje", "cf_viaje", "$"), ("Subtotal costos variables del viaje", "cv_viaje", "$"),
             ("Comisiones + factor prestacional", "o_comisiones", "$"), ("Factor administrativo", "o_admin", "$"), ("Retefuente + ICA", "o_rete", "$"),
             ("Peajes del viaje", "v_peajes", "$"), ("Costo hora (publicado)", "costo_hora", "$/h"),
             ("Costo movilización carga (publicado)", "costo_mov_carga", "$"), ("Costo total del viaje (publicado)", "costo_total_viaje", "$"),
             ("Distancia total", "dist_total_km", "km")]
    R = {}
    for i, (n, key, un) in enumerate(filas):
        r = 4 + i
        R[key] = r
        C.c(ws, r, 1, n)
        fm = "0.0%" if key == "tasa_imprevistos" else ("#,##0.00" if un in ("$/km", "h", "km") else '"$"#,##0.00')
        C.c(ws, r, 2, s1[key], C.ENTRADA, fm); C.c(ws, r, 3, s2[key], C.ENTRADA, fm); C.c(ws, r, 4, un)
    r = 4 + len(filas) + 1
    C.cab(ws, r, ["Terreno", "km " + s1["ruta"], "km " + s2["ruta"], "Velocidad (km/h)", "Rendimiento (km/galón)"], col0=1)
    ws.column_dimensions["E"].width = 18
    RT = {}
    for i, t in enumerate(C.TERRENOS):
        rr = r + 1 + i
        RT[t] = rr
        C.c(ws, rr, 1, t); C.c(ws, rr, 2, s1["dist_km"][t], C.ENTRADA, "0.00"); C.c(ws, rr, 3, s2["dist_km"][t], C.ENTRADA, "0.00")
        C.c(ws, rr, 4, s1["vel_kmh"][t], C.ENTRADA, "0.00"); C.c(ws, rr, 5, s1["rend_kmgal"][t], C.ENTRADA, "0.00")
    C.nota(ws, r + 7, "Todas las cifras en azul se leen automáticamente del PDF de la consulta (script, sin transcripción manual). "
                      "El SICE-TAC publica en formato numérico en-US; aquí se muestran en es-CO.", ancho=5)
    # ---- Unitarios
    u = wb.create_sheet("Unitarios"); u.sheet_view.showGridLines = False
    C.tit(u, "Tabla 2. Costos unitarios del tractocamión 3S3 (pesos de septiembre de 2026)")
    C.cab(u, 3, ["Unidad de costo", "Valor", "Unidad", "Cómo se obtiene", "Marca"], [44, 16, 10, 64, 10])
    U = {}
    lista = [("Costo fijo por hora de vehículo", f"=Insumos!B{R['cf_mes']}/Insumos!B{R['horas_habiles_mes']}", "$/h", "Costos fijos del mes / horas hábiles del mes", '"$"#,##0.00', "cf_hora"),
             ("Costo variable por km sin combustible", f"=SUM(Insumos!B{R['km_llantas']}:Insumos!B{R['km_lavado']})*(1+Insumos!B{R['tasa_imprevistos']})", "$/km",
              "(Llantas + lubricantes + filtros + mantenimiento + lavado) × (1 + imprevistos)", '"$"#,##0.00', "cv"),
             ("Factor de otros costos f", f"=(Insumos!B{R['o_comisiones']}+Insumos!B{R['o_admin']}+Insumos!B{R['o_rete']})/(Insumos!B{R['cf_viaje']}+Insumos!B{R['cv_viaje']})",
              "—", "(Comisiones + administrativo + retefuente e ICA) / (fijos + variables del viaje)", "0.0000%", "f"),
             ("Costo por hora con otros costos", "=B4*(1+B6)", "$/h", "Costo fijo por hora × (1 + f). Debe igualar el «costo hora» publicado", '"$"#,##0.00', "ch"),
             ("Control: costo hora publicado", f"=Insumos!B{R['costo_hora']}", "$/h", "SICE-TAC (resumen de la consulta)", '"$"#,##0.00', "chp"),
             ("Diferencia relativa", "=B7/B8-1", "—", "Criterio: |dif.| ≤ 0,1 %", "0.000%", "d")]
    for i, (n, fml, un, como, fm, key) in enumerate(lista):
        r = 4 + i
        U[key] = r
        C.c(u, r, 1, n, C.NEGRA if key in ("cf_hora", "cv", "f", "ch") else C.NORMAL); C.c(u, r, 2, fml, fmt=fm, fill=C.FILL_RES if key in ("cf_hora", "cv", "ch") else None)
        C.c(u, r, 3, un); C.c(u, r, 4, como, wrap=True); C.c(u, r, 5, "[CP]" if key != "chp" else "[F]")
    r = 12
    u.cell(r, 1, "Tabla 3. Combustible por kilómetro y tipo de terreno").font = C.Font(name=C.F, size=13, bold=True, color=C.AZUL)
    C.cab(u, r + 1, ["Terreno", "Combustible ($/km)", "Unidad", "Cómo se obtiene", "Marca"])
    for i, t in enumerate(C.TERRENOS):
        rr = r + 2 + i
        U["comb_" + t] = rr
        C.c(u, rr, 1, t); C.c(u, rr, 2, f"=Insumos!B{R['acpm']}/Insumos!E{RT[t]}", fmt='"$"#,##0.00', fill=C.FILL_RES)
        C.c(u, rr, 3, "$/km"); C.c(u, rr, 4, "Valor ACPM / rendimiento del terreno"); C.c(u, rr, 5, "[CP]")
    r = 20
    u.cell(r, 1, "Tabla 4. Desglose del factor de otros costos").font = C.Font(name=C.F, size=13, bold=True, color=C.AZUL)
    C.cab(u, r + 1, ["Componente", "Tasa sobre fijos + variables", "", "Rótulo en el SICE-TAC", "Marca"])
    for i, (n, key, rot) in enumerate([("Comisiones + factor prestacional", "o_comisiones", "Comisiones + Factor Prestacional"),
                                       ("Factor administrativo", "o_admin", "Factor Administrativo (5 %)"), ("Retefuente + ICA", "o_rete", "Retefuente + ICA (3,5 % + 0,3 %) = 3,8 %")]):
        rr = r + 2 + i
        C.c(u, rr, 1, n); C.c(u, rr, 2, f"=Insumos!B{R[key]}/(Insumos!B{R['cf_viaje']}+Insumos!B{R['cv_viaje']})", fmt="0.000%"); C.c(u, rr, 4, rot); C.c(u, rr, 5, "[CP]")
    C.nota(u, r + 6, "El SICE-TAC reparte los costos fijos del mes entre los viajes por horas (factor de viajes = horas hábiles / horas de recorrido); "
                     "por eso el costo fijo es un costo por HORA de vehículo, que vale tanto conduciendo como esperando o viajando sobre el tren. "
                     "Horas hábiles: 230 h/mes, vigentes desde el 15 jul 2026 (el SICE-TAC las bajó de 238 h por la Ley 2101 de 2021, según las «fechas de actualización» de la propia consulta); el protocolo de la Res. 20243040057465 de 2024 partía de 288 h.", ancho=5)
    # ---- Verificación
    v = wb.create_sheet("Verificación"); v.sheet_view.showGridLines = False
    C.tit(v, "Tabla 5. Reproducción del costo SICE-TAC con los costos unitarios (criterio: diferencia ≤ 0,1 %)")
    C.cab(v, 3, ["Concepto", s1["ruta"], s2["ruta"], "Cómo se obtiene"], [40, 18, 18, 70])
    col = {"B": "B", "C": "C"}
    lin = [("Horas de viaje", lambda X: "=" + "+".join(f"Insumos!{X}{RT[t]}/Insumos!D{RT[t]}" for t in C.TERRENOS), "0.0000", "Σ km / velocidad por terreno"),
           ("Combustible del viaje", lambda X: "=" + "+".join(f"Insumos!{X}{RT[t]}*Unitarios!B{U['comb_' + t]}" for t in C.TERRENOS), '"$"#,##0', "Σ km × combustible por km"),
           ("Variables sin combustible", lambda X: f"=Insumos!{X}{R['dist_total_km']}*Unitarios!B{U['cv']}", '"$"#,##0', "km totales × costo variable por km"),
           ("Peajes", lambda X: f"=Insumos!{X}{R['v_peajes']}", '"$"#,##0', "Listado oficial de peajes (Tabla 6)"),
           ("Costo de movilización reproducido", lambda X: f"=(Unitarios!$B$4*{X}4+{X}5+{X}6+{X}7)*(1+Unitarios!$B$6)", '"$"#,##0', "(costo fijo/h × horas + variables) × (1 + f)"),
           ("Costo de movilización publicado", lambda X: f"=Insumos!{X}{R['costo_mov_carga']}", '"$"#,##0', "SICE-TAC [F]"),
           ("Diferencia relativa", lambda X: f"={X}8/{X}9-1", "0.000%", ""),
           ("Costo total del viaje reproducido", lambda X: f"=(Unitarios!$B$4*({X}4+Insumos!{X}{R['horas_espera']})+{X}5+{X}6+{X}7)*(1+Unitarios!$B$6)", '"$"#,##0', "Incluye las 4 h logísticas mínimas"),
           ("Costo total del viaje publicado", lambda X: f"=Insumos!{X}{R['costo_total_viaje']}", '"$"#,##0', "SICE-TAC [F]"),
           ("Diferencia relativa", lambda X: f"={X}11/{X}12-1", "0.000%", "")]
    for i, (n, f, fm, como) in enumerate(lin):
        r = 4 + i
        C.c(v, r, 1, n, C.NEGRA if "Diferencia" in n else C.NORMAL)
        for X in ("B", "C"):
            C.c(v, r, 2 if X == "B" else 3, f(X), fmt=fm, fill=C.FILL_RES if "Diferencia" in n else None)
        C.c(v, r, 4, como, wrap=True)
    C.nota(v, 15, "Si alguna diferencia supera el 0,1 %, el script se detiene antes de escribir este archivo. Resultado de la corrida: " +
           "; ".join(f"{S['sentidos'][k]['ruta']}: movilización {100 * S['sentidos'][k]['error_rel']['movilizacion']:+.4f} %, total {100 * S['sentidos'][k]['error_rel']['total']:+.4f} %".replace(".", ",")
                     for k in C.RUTAS) + ".", ancho=4)
    # ---- Peajes
    pj = wb.create_sheet("Peajes"); pj.sheet_view.showGridLines = False
    C.tit(pj, "Tabla 6. Peajes de la Ruta 40 entre Ibagué y Calarcá, categoría V (tractocamión 3S3)")
    C.cab(pj, 3, ["Sentido", "Peaje", "Tarifa ($)", "Marca"], [22, 34, 14, 40])
    r = 4
    for k in C.RUTAS:
        for n, val in S["peajes"][k]:
            C.c(pj, r, 1, S["sentidos"][k]["ruta"]); C.c(pj, r, 2, n); C.c(pj, r, 3, val, C.ENTRADA, '"$"#,##0'); C.c(pj, r, 4, "[F] listado SICE-TAC 01-09-2026")
            r += 1
    C.nota(pj, r + 1, "El peaje de Cocora solo se cobra en sentido Ibagué → Cajamarca; por eso el viaje de ida paga 106.700 $ y el de regreso 42.900 $.", ancho=4)
    C.hoja_fuentes(wb, ["SICETAC_IBG_CAL", "SICETAC_CAL_IBG", "SICETAC_PEAJ", "RES2024", "RES2026"], 7)
    for w in wb.worksheets:
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(SALIDA_X)


if __name__ == "__main__":
    main()
